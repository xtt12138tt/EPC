"""
Playwright 浏览器自动化 - 填写 EPC 报销表单
"""
import asyncio
import time
import json
import re
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from playwright.async_api import async_playwright, Page, Locator
from pypinyin import lazy_pinyin

from config import (
    EPC_URL, EPC_USER, EPC_PASS, BROWSER_CHANNEL, CHROME_USER_DATA,
    TIMEOUT, SLOW_MO, REIMBURSEMENT_TYPE_MAP, FANBAOTYPE_MAP,
)


# ─────────────────────────────────────────────
# Ant Design 组件操作工具函数
# ─────────────────────────────────────────────

async def ant_select(page: Page, placeholder: str, value: str, container_sel: str = "body"):
    """点击 Ant Design 下拉选择器并选择选项"""
    # 找到含对应 placeholder 的 ant-select 容器
    selector = f'{container_sel} .ant-select-selection__placeholder:text("{placeholder}")'
    alt_selector = f'{container_sel} .ant-select-selection-selected-value'

    # 点击触发下拉
    select_el = page.locator(
        f'{container_sel} .ant-select'
    ).filter(has=page.locator(f':text("{placeholder}")').or_(
        page.locator(f'[placeholder="{placeholder}"]')
    )).first

    # 更可靠的方式：找最近的 .ant-select-selection
    select_wrapper = page.locator(
        f'.ant-select-selection'
    ).filter(has=page.locator(f'.ant-select-selection__placeholder:has-text("{placeholder}")')).first

    await select_wrapper.click()
    await page.wait_for_selector(".ant-select-dropdown:not(.ant-select-dropdown-hidden)", timeout=TIMEOUT)
    option = page.locator(f'.ant-select-dropdown-menu-item:has-text("{value}")').first
    await option.wait_for(timeout=TIMEOUT)
    await option.click()
    await page.wait_for_timeout(200)


async def ant_select_search(page: Page, placeholder: str, search_text: str, container_sel: str = "body"):
    """可搜索的 Ant Design 下拉（输入文字后选择）"""
    # 找到 select wrapper
    wrapper = page.locator(
        '.ant-select-selection'
    ).filter(
        has=page.locator(f'.ant-select-selection__placeholder:has-text("{placeholder}")')
    ).first
    if await wrapper.count() == 0:
        # 兜底：EPC 改版后 placeholder 文案可能微调，取核心词模糊匹配
        keyword = placeholder.split("】")[-1].split("的")[-1].strip()
        if keyword and keyword != placeholder:
            wrapper = page.locator(
                '.ant-select-selection'
            ).filter(
                has=page.locator(f'.ant-select-selection__placeholder:has-text("{keyword}")')
            ).first
        if await wrapper.count() == 0:
            raise TimeoutError(f"未找到下拉框：{placeholder}")

    await wrapper.click()
    await page.wait_for_timeout(300)

    # 输入搜索
    await page.keyboard.type(search_text, delay=80)

    # 等下拉菜单出现并有可选项
    await page.wait_for_selector(
        ".ant-select-dropdown:not(.ant-select-dropdown-hidden) .ant-select-dropdown-menu-item",
        timeout=TIMEOUT
    )
    await page.wait_for_timeout(300)

    # 点击第一个非禁用选项
    option = page.locator(
        '.ant-select-dropdown:not(.ant-select-dropdown-hidden) '
        '.ant-select-dropdown-menu-item:not(.ant-select-dropdown-menu-item-disabled)'
    ).first
    await option.click()
    await page.wait_for_timeout(300)


async def react_fill_number(inp: Locator, value) -> None:
    """向 React 受控的 ant-input-number 写值，触发 onChange"""
    await inp.click()
    await inp.evaluate("el => { el.select(); }")
    await inp.fill("")
    await inp.press_sequentially(str(value), delay=30)
    await inp.press("Tab")
    await inp.evaluate(
        "el => { el.dispatchEvent(new Event('input', {bubbles:true})); "
        "el.dispatchEvent(new Event('change', {bubbles:true})); }"
    )


async def ant_number_input(page: Page, placeholder: str, value, container_sel: str = "body"):
    """填写 Ant Design 数字输入框"""
    inp = page.locator(
        f'.ant-input-number-input[placeholder="{placeholder}"]'
    ).first
    await react_fill_number(inp, value)
    await page.wait_for_timeout(200)


def _normalize_select_option(value: str) -> str:
    """统一 EPC 下拉文案，兼容小时、单位、标点等小差异。"""
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", "", text)
    text = text.replace("～", "~").replace("－", "-").replace("—", "-")
    text = text.replace("-", "~").replace("：", ":")
    text = re.sub(r"(\d+(?:\.\d+)?)h", r"\1小时", text)
    text = text.replace("hours", "小时").replace("hour", "小时")
    text = text.replace("mins", "min").replace("minutes", "min").replace("minute", "min")
    text = text.replace("小时以内", "小时内").replace("min以内", "min内")
    text = text.replace("样本", "单位").replace("个", "单位")
    text = text.replace("命中率:", "命中率")
    return text


def _duration_bounds_minutes(value: str):
    """解析下拉文案中的时长区间，统一换算为分钟。"""
    text = _normalize_select_option(value)
    range_match = re.search(r"(\d+(?:\.\d+)?)~(\d+(?:\.\d+)?)(小时|min)", text)
    if range_match:
        start, end, unit = range_match.groups()
        multiplier = 60 if unit == "小时" else 1
        return float(start) * multiplier, float(end) * multiplier, False

    upper_match = re.search(r"(?:<|≤)?(\d+(?:\.\d+)?)(小时|min)内", text)
    if upper_match:
        end, unit = upper_match.groups()
        multiplier = 60 if unit == "小时" else 1
        return 0.0, float(end) * multiplier, True

    less_match = re.search(r"<(\d+(?:\.\d+)?)(小时|min)", text)
    if less_match:
        end, unit = less_match.groups()
        multiplier = 60 if unit == "小时" else 1
        return 0.0, float(end) * multiplier, True

    return None


def _duration_match_score(source: str, option: str):
    """为“原始兼职难度 → EPC 选项”计算区间匹配分，无法判断时返回 None。"""
    source_bounds = _duration_bounds_minutes(source)
    option_bounds = _duration_bounds_minutes(option)
    if not source_bounds or not option_bounds:
        return None

    source_start, source_end, source_is_upper = source_bounds
    option_start, option_end, _ = option_bounds
    if option_start <= source_start and source_end <= option_end:
        # EPC 选项完整覆盖原始区间：范围越窄，匹配越精确。
        return 10_000 - (option_end - option_start)
    if source_is_upper and abs(option_end - source_end) < 0.001:
        # “30min以内”跨多个平台档位时，按其上限落到“10~30min”这一档。
        return 5_000 - (option_end - option_start)
    return None


async def ant_select_in_row(row: Locator, col_index: int, value: str, page: Page):
    """在表格行选择下拉项，并兼容 EPC 文案的常见差异。"""
    cells = row.locator("td")
    cell = cells.nth(col_index)
    wrapper = cell.locator(".ant-select-selection").first
    await wrapper.click()
    await page.wait_for_selector(
        ".ant-select-dropdown:not(.ant-select-dropdown-hidden)", timeout=TIMEOUT
    )
    items = page.locator(
        '.ant-select-dropdown:not(.ant-select-dropdown-hidden) '
        '.ant-select-dropdown-menu-item'
    )
    target = _normalize_select_option(value)
    options = []
    option_rows = []
    for index in range(await items.count()):
        option = items.nth(index)
        label = (await option.inner_text()).strip()
        options.append(label)
        option_rows.append((option, label))
        if label == value or _normalize_select_option(label) == target:
            await option.click()
            await page.wait_for_timeout(200)
            return label

    duration_candidates = [
        (score, option, label)
        for option, label in option_rows
        if (score := _duration_match_score(value, label)) is not None
    ]
    if duration_candidates:
        _, option, label = max(duration_candidates, key=lambda item: item[0])
        await option.click()
        await page.wait_for_timeout(200)
        return label
    raise ValueError(f"下拉未找到“{value}”，当前可选项：{options}")


async def ant_number_in_row(row: Locator, placeholder: str, value, quiet: bool = False) -> bool:
    """在表格行中按 placeholder 填数字输入框（支持精确 / 包含匹配）"""
    # 先精确匹配
    inp = row.locator(f'.ant-input-number-input[placeholder="{placeholder}"]').first
    if await inp.count() == 0:
        # 再做包含匹配（页面 placeholder 可能有括号/空格差异）
        key = placeholder.replace("请输入", "").replace("（元）", "").replace("(元)", "").strip()
        inp = row.locator(f'.ant-input-number-input[placeholder*="{key}"]').first
    if await inp.count() > 0:
        await react_fill_number(inp, value)
        return True
    else:
        # 打印该行所有 number input 的 placeholder，帮助调试
        all_inputs = row.locator('.ant-input-number-input')
        cnt = await all_inputs.count()
        placeholders = [await all_inputs.nth(j).get_attribute("placeholder") for j in range(cnt)]
        if not quiet:
            print(f"    ⚠ 未找到 placeholder='{placeholder}' 的数字框，该行实际 placeholders: {placeholders}")
        return False


# ─────────────────────────────────────────────
# 登录
# ─────────────────────────────────────────────

async def login(page: Page):
    """登录 EPC 平台 - 自动检测登录态或等待用户手动 SSO"""
    print("→ 打开 EPC 平台...")
    await page.goto(EPC_URL, wait_until="domcontentloaded", timeout=30000)

    # 轮询 10 秒：表单出现 → 已登录；跳到 SSO → 需要手动登录
    for _ in range(20):  # 20 × 500ms = 10s
        await page.wait_for_timeout(500)
        if await page.locator(".ant-select-selection").count() > 0:
            print("✓ 已检测到登录态，表单已加载")
            return
        if "login.netease.com" in page.url:
            break  # SSO 跳转发生了
    else:
        # 10s 后既没表单也没跳 SSO —— 再等15s
        try:
            await page.wait_for_selector(".ant-select-selection", timeout=15000)
            print("✓ 表单已加载")
            return
        except Exception:
            pass  # 还是没出来，走 SSO 流程

    # ── 需要手动 SSO 登录 ──
    print("\n" + "="*55)
    print("⚠  请在弹出的浏览器窗口里完成网易 SSO 登录")
    print("   登录成功后脚本会自动继续（最多等待 5 分钟）")
    print("="*55)

    # 轮询等待：每秒检测一次，直到表单出现（不依赖导航事件）
    for i in range(300):  # 300 × 1s = 5 分钟
        await page.wait_for_timeout(1000)
        try:
            cur_url = page.url
            if "eplayer.nie.netease.com" in cur_url and "login.netease.com" not in cur_url:
                cnt = await page.locator(".ant-select-selection").count()
                if cnt > 0:
                    print(f"✓ SSO 登录完成，表单已加载（等待了约 {i+1} 秒）")
                    return
                elif "#/login" not in cur_url:
                    # 已回到 EPC 但表单还没渲染，导航一次
                    break
        except Exception:
            pass

    # 登录后重新导航到报销页
    await page.goto(EPC_URL, wait_until="domcontentloaded", timeout=30000)
    await page.wait_for_selector(".ant-select-selection", timeout=30000)
    print("✓ 报销表单已加载")


async def _ask_runtime(runtime: Any, prompt: str, status: str) -> str:
    """兼容旧命令行与新批次 Worker 的确认输入。"""
    if runtime is not None:
        return await runtime.ask(prompt, status)
    return input(prompt)


async def _checkpoint(runtime: Any, step: str) -> None:
    """批次 Worker 的安全暂停点；旧命令行模式下为空操作。"""
    if runtime is not None:
        await runtime.checkpoint(step)


async def wait_and_click_next(
    page: Page,
    runtime: Any,
    stage: str,
    next_page_marker: str,
) -> bool:
    """确认后进入下一页，兼容脚本自动点击和用户手动点击两种方式。

    EPC 会在必填字段、上传文件或异步校验尚未完成时禁用按钮。这里不再
    因固定 30 秒超时而失败；用户补齐页面字段后，按钮恢复可用会自动继续。
    等待过程不重复写日志，避免执行面板被“已等待”信息刷屏。

    若用户已经手动点了“下一步”，先检测下一页标志并直接继续，绝不对
    已切换的页面再次点击。若标签页/浏览器上下文被关闭，返回 False 由
    调用方标记为待修正，而不是抛出 TargetClosedError。
    """
    click_sent = False
    while True:
        await _checkpoint(runtime, f"等待 {stage} 下一步按钮可用")
        if page.is_closed():
            if runtime is not None:
                runtime.log(f"⚠ {stage}：EPC 标签页已关闭或自动化连接已中断")
            return False
        try:
            if await page.locator(next_page_marker).count() > 0:
                if runtime is not None:
                    runtime.log(f"✓ {stage}：已检测到下一页（脚本或手动操作均可）")
                return True
            next_button = page.locator('button.ant-btn-primary:has-text("下一步")').last
            enabled = await next_button.is_enabled()
            visible = await next_button.is_visible()
        except Exception:
            enabled = False
            visible = False

        if visible and enabled and not click_sent:
            try:
                await next_button.click(timeout=1_000)
                click_sent = True
                if runtime is not None:
                    runtime.log(f"✓ {stage}：下一步按钮已可用，已发起进入下一页")
            except Exception:
                # EPC 可能在本次渲染周期内再次禁用；继续观察即可。
                pass

        await page.wait_for_timeout(1_000)


def _page1_payload(data: dict) -> tuple[dict, str]:
    """将标准单据 JSON 转为 fill_page1 使用的旧结构。"""
    from expense_note import build_standard_expense_note

    expense_note = build_standard_expense_note(data)
    return {
        "overview": {
            "project_id": data.get("project_id"),
            "fanbao_series": data.get("fanbao_series"),
            "fanbao_type": data.get("fanbao_type"),
            "cost_center": data.get("cost_center"),
            "screenshot_required": data.get("screenshot_required", True),
            "screenshot_mode": data.get("screenshot_mode"),
            "screenshot_path": data.get("screenshot_path"),
            "reimbursement_types": data.get("reimbursement_types", []),
            "expense_note_manual": expense_note,
        },
        "gift_common": data.get("gift_common"),
        "gift_special": data.get("gift_special") or [],
        "questionnaire": data.get("questionnaire") or [],
        "parttime": data.get("parttime"),
        "other": data.get("other"),
        "payee_rules": data.get("payee_rules", {}),
        "split_note": data.get("split_note", {}),
    }, expense_note



# ─────────────────────────────────────────────
# 第1页：报销内容
# ─────────────────────────────────────────────

async def fill_page1(page: Page, data: dict, expense_note: str, runtime: Any = None):
    """填写第1页所有字段"""
    ov = data["overview"]
    print("→ 第1页：填写报销内容...")

    # 1. 发包序列
    if ov.get("fanbao_series"):
        print(f"  发包序列 = {ov['fanbao_series']}")
        await ant_select(page, "请选择发包序列", ov["fanbao_series"])
        await _checkpoint(runtime, "第1页发包序列填写后")

    # 2. 发包类型（radio）
    if ov.get("fanbao_type"):
        radio_val = FANBAOTYPE_MAP.get(ov["fanbao_type"], "person")
        print(f"  发包类型 = {ov['fanbao_type']} ({radio_val})")
        await page.locator(f'input[type="radio"][value="{radio_val}"]').click()
        await page.wait_for_timeout(200)
        await _checkpoint(runtime, "第1页发包类型填写后")

    # 3. 报销项目（可搜索下拉）
    if ov.get("project_id"):
        print(f"  报销项目 = {ov['project_id']}")
        await ant_select_search(
            page,
            "可关联【执行中、已完结】且有采购单的研究项目",
            ov["project_id"]
        )
        await _checkpoint(runtime, "第1页报销项目填写后")

    # 4. 成本归属（可搜索下拉）
    if ov.get("cost_center"):
        print(f"  成本归属 = {ov['cost_center']}")
        await ant_select_search(page, "请确认报销产品或部门", ov["cost_center"])
        await page.wait_for_timeout(1000)  # 等费用主体/预算负责人联动
        await _checkpoint(runtime, "第1页成本归属填写后")

    # 5. 费用说明（textarea）
    print(f"  费用说明（长度 {len(expense_note)} 字）")
    textarea = page.locator("textarea.ant-input").first
    await textarea.click()
    await textarea.fill(expense_note)
    await _checkpoint(runtime, "第1页费用说明填写后")

    # 6. 产品确认截图（文件上传）
    # screenshot_mode:
    #   required   - 必须上传；没有文件或上传控件时给出警告
    #   if_present - 当前 EPC 页存在上传控件时上传，否则正常跳过
    #   skip       - 不处理截图
    screenshot_required = ov.get("screenshot_required", True)
    screenshot_mode = ov.get("screenshot_mode") or ("required" if screenshot_required else "skip")

    if screenshot_mode == "skip":
        print("  产品确认截图：本单无需上传，跳过")
    else:
        screenshot_paths = [
            p.strip()
            for p in (ov.get("screenshot_path") or "").split(";")
            if p.strip()
        ]
        file_inputs = page.locator('input[type="file"]')
        upload_count = await file_inputs.count()
        if screenshot_mode == "if_present" and upload_count == 0:
            print("  产品确认截图：当前 EPC 页面未出现上传控件，按条件规则跳过")
        elif screenshot_paths:
            valid = [p for p in screenshot_paths if Path(p).exists()]
            if valid:
                print(f"  上传截图 {len(valid)} 个文件")
                file_input = file_inputs.first
                await file_input.set_input_files(valid)
                await page.wait_for_timeout(500)
            else:
                print(f"  ⚠ 截图文件不存在: {screenshot_paths}，请在浏览器中手动上传")
        else:
            print("  ⚠ 需要上传产品确认截图，请在浏览器中手动上传后再确认第1页")
    await _checkpoint(runtime, "第1页产品截图处理后")


    # 7. 报销明细复选框
    reimbursement_types = ov.get("reimbursement_types", [])
    print(f"  报销明细 = {reimbursement_types}")
    for t in reimbursement_types:
        val = REIMBURSEMENT_TYPE_MAP.get(t)
        if val:
            cb = page.locator(f'input[type="checkbox"][value="{val}"]')
            if not await cb.is_checked():
                await cb.click()
            await page.wait_for_timeout(300)
    await _checkpoint(runtime, "第1页报销明细勾选后")

    # 8. 填写各明细子表
    if any(t in REIMBURSEMENT_TYPE_MAP for t in reimbursement_types
           if REIMBURSEMENT_TYPE_MAP.get(t) == "domesticCommon"):
        await fill_gift_common(page, data)
        await _checkpoint(runtime, "第1页常规礼金填写后")

    if any(REIMBURSEMENT_TYPE_MAP.get(t) == "domestic" for t in reimbursement_types):
        await fill_gift_special(page, data)
        await _checkpoint(runtime, "第1页特殊礼金填写后")

    if any(REIMBURSEMENT_TYPE_MAP.get(t) == "questionnaire" for t in reimbursement_types):
        await fill_questionnaire(page, data)
        await _checkpoint(runtime, "第1页问卷明细填写后")

    if any(REIMBURSEMENT_TYPE_MAP.get(t) == "partTime" for t in reimbursement_types):
        await fill_parttime(page, data)
        await _checkpoint(runtime, "第1页兼职明细填写后")

    if any(REIMBURSEMENT_TYPE_MAP.get(t) == "other" for t in reimbursement_types):
        await fill_other(page, data)
        await _checkpoint(runtime, "第1页其他费用填写后")

    print("  ✓ 第1页填写完成")


async def fill_gift_common(page: Page, data: dict):
    """填写国内玩家礼金（常规）子表"""
    gc = data.get("gift_common", {})
    rows = gc.get("rows", [])
    if not rows:
        return

    print(f"  → 礼金(常规): 总样本量={gc.get('total_sample')}, 行数={len(rows)}")

    # 等子表单展开
    await page.wait_for_timeout(800)

    # 定位总样本量：直接用 placeholder 找，不依赖 #domesticCommon
    total_sample_inp = page.locator('input.ant-input-number-input[placeholder="请输入总样本量"]')
    ts_count = await total_sample_inp.count()
    print(f"    总样本量输入框: {ts_count} 个")
    if ts_count > 0 and gc.get("total_sample") is not None:
        await react_fill_number(total_sample_inp.first, gc["total_sample"])
        await page.wait_for_timeout(500)  # 等 React 渲染子表行

    # 找展开的子表：包含"测试形式"列标题的表格
    # 用标题文字定位所在的 table，再找 tbody
    gift_common_table = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("测试形式")')
    ).first
    tc = await gift_common_table.count()
    print(f"    礼金(常规)子表: {tc} 个")

    if tc == 0:
        print("    ⚠ 未找到礼金(常规)子表，跳过行填写")
        return

    tbody = gift_common_table.locator(".ant-table-tbody")

    first_row_defaults = {}  # 记录第1行的 测试形式/连续周期，作为后续行的 fallback

    for i, row_data in enumerate(rows):
        # 后续行若缺少级联前置字段，自动继承第1行的值
        if i == 0:
            first_row_defaults = {k: v for k, v in row_data.items()
                                   if k in ("测试形式", "连续周期", "样本稀缺性") and v}
            effective = row_data
        else:
            effective = {**first_row_defaults, **row_data}

        # 需要加行时点+号（用多个候选选择器）
        if i > 0:
            add_btn = (
                gift_common_table.locator('[class*="addIcon"]').last
                if await gift_common_table.locator('[class*="addIcon"]').count() > 0
                else page.locator('.anticon-plus-circle').last
            )
            await add_btn.click()
            # 等待新行出现
            await page.wait_for_timeout(600)
            await page.wait_for_selector(
                f'.ant-table-tbody tr.ant-table-row:nth-child({i+1})',
                timeout=5000
            )

        table_rows = tbody.locator("tr.ant-table-row")
        row = table_rows.nth(i)
        print(f"    填第{i+1}行: {effective}")

        # 列顺序：测试形式(0), 连续周期(1), 礼金小类(2), 样本稀缺性(3)
        # 顺序很重要：礼金小类 会影响后续列结构，选完必须等 UI 稳定
        col_map = [
            ("测试形式", 0),
            ("连续周期", 1),
            ("礼金小类", 2),
            ("样本稀缺性", 3),
        ]
        for field, col_idx in col_map:
            v = effective.get(field)
            if not (v and str(v).strip() and str(v).strip() != "-"):
                continue
            for attempt in range(2):
                try:
                    await ant_select_in_row(row, col_idx, str(v).strip(), page)
                    break
                except Exception as e:
                    print(f"    ⚠ {field} 填写异常(第{attempt+1}次): {e}")
                    await page.keyboard.press("Escape")
                    await page.wait_for_timeout(400)
            # 选完礼金小类：UI 会 rerender（转介费只显示单价，无测试时长等），等它稳定
            if field == "礼金小类":
                await page.wait_for_timeout(600)

        # 测试时长（数字输入，第4列）
        duration = effective.get("测试时长(小时)")
        if duration is not None:
            await ant_number_in_row(row, "请输入测试时长（小时）", duration)

        # 招募周期（仅礼金小类=应急礼金时出现）
        recruit_cycle = effective.get("招募周期")
        if recruit_cycle and str(recruit_cycle).strip() not in ("", "-"):
            cells = row.locator("td")
            recruit_cell = cells.nth(5)
            sel = recruit_cell.locator(".ant-select-selection").first
            if await sel.count() > 0:
                await sel.click()
                await page.wait_for_timeout(300)
                opt = page.locator(
                    f'.ant-select-dropdown-menu-item:has-text("{recruit_cycle}")'
                ).first
                await opt.click()

        # 样本量 / 总金额 / 单价（转介费用单价，其他用总金额）
        sample = effective.get("样本量")
        if sample is not None:
            await ant_number_in_row(row, "请输入样本量", sample)

        unit_price = effective.get("单价(元)") or effective.get("单价（元）") or effective.get("单价")
        if unit_price is not None:
            # 尝试多个 placeholder
            for ph in ("请输入单价（元）", "请输入单价(元)", "请输入单价"):
                inp = row.locator(f'input[placeholder="{ph}"]').first
                if await inp.count() > 0:
                    await inp.triple_click()
                    await inp.fill(str(unit_price))
                    break

        total = effective.get("总金额(元)")
        if total is not None:
            await ant_number_in_row(row, "请输入总金额（元）", total)

        await page.wait_for_timeout(200)


async def fill_gift_special(page: Page, data: dict):
    """填写国内玩家礼金（小众/特殊）子表"""
    rows = data.get("gift_special", [])
    if not rows:
        return

    print(f"  → 礼金(小众/特殊): 行数={len(rows)}")
    container = page.locator("#domestic")
    tbody = container.locator(".ant-table-tbody")

    for i, row_data in enumerate(rows):
        if i > 0:
            await container.locator(".addIcon___1KN9Y").last.click()
            await page.wait_for_timeout(500)

        table_rows = tbody.locator("tr.ant-table-row")
        row = table_rows.nth(i)

        col_map = {"样本类型": 0, "测试形式": 1, "地区": 2, "难度等级": 3}
        for field, col_idx in col_map.items():
            v = row_data.get(field)
            if v and str(v).strip() and str(v).strip() != "-":
                await ant_select_in_row(row, col_idx, str(v).strip(), page)

        for placeholder, field in [
            ("请输入测试时长", "测试时长"),
            ("请输入样本量", "测试样本量"),
            ("请输入总金额（元）", "测试总金额"),
        ]:
            v = row_data.get(field)
            if v is not None:
                inp = row.locator(f'input[placeholder="{placeholder}"]').first
                if await inp.count() > 0:
                    await inp.triple_click()
                    await inp.fill(str(v))

        await page.wait_for_timeout(200)


async def fill_questionnaire(page: Page, data: dict):
    """填写国内问卷调研子表"""
    qt = data.get("questionnaire", {})
    rows = qt.get("rows", []) if isinstance(qt, dict) else (qt or [])
    if not rows:
        return

    print(f"  → 问卷调研: 行数={len(rows)}")
    await page.wait_for_timeout(500)

    questionnaire_table = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("投放渠道")')
    ).first
    if await questionnaire_table.count() == 0:
        print("    ⚠ 未找到问卷调研子表")
        return
    tbody = questionnaire_table.locator(".ant-table-tbody")

    for i, row_data in enumerate(rows):
        if i > 0:
            await container.locator(".addIcon___1KN9Y").last.click()
            await page.wait_for_timeout(500)

        table_rows = tbody.locator("tr.ant-table-row")
        row = table_rows.nth(i)

        col_map = {"国家": 0, "投放渠道": 1, "题目数量": 2, "难度/调研方式": 3}
        for field, col_idx in col_map.items():
            v = row_data.get(field)
            if v and str(v).strip() and str(v).strip() != "-":
                await ant_select_in_row(row, col_idx, str(v).strip(), page)

        sample = row_data.get("样本量")
        if sample is not None:
            await ant_number_in_row(row, "请输入样本量", sample)

        total = row_data.get("总金额(元)")
        if total is not None:
            await ant_number_in_row(row, "请输入总金额（元）", total)

        await page.wait_for_timeout(200)


async def fill_parttime(page: Page, data: dict):
    """填写国内兼职子表"""
    pt = data.get("parttime", {})
    rows = pt.get("rows", []) if isinstance(pt, dict) else (pt or [])
    if not rows:
        return

    def normalize_difficulty(v: str) -> str:
        """将兼职难度映射为 EPC 标准文案。"""
        v = str(v or "").strip().replace("-", "~").replace("～", "~")
        _hour_map = {
            "30min以内/样本": "30min以内/样本",
            "30~60min/样本": "30~60min/样本",
            "10min以内/样本": "10min以内/样本",
            "10~30min/样本": "10~30min/样本",
            "60~120min/样本": "60~120min/样本",
            "10~20min/样本": "10~20min/样本",
            "20~30min/样本": "20~30min/样本",
            "命中率：>10%": "命中率：>10%",
            "命中率:>10%": "命中率：>10%",
            "命中率：6-10%": "命中率：6-10%",
            "命中率:6-10%": "命中率：6-10%",
            "命中率：4-6%": "命中率：4-6%",
            "命中率:4-6%": "命中率：4-6%",
            "命中率：2-4%": "命中率：2-4%",
            "命中率:2-4%": "命中率：2-4%",
            "命中率：<2%": "命中率：<2%",
            "命中率:<2%": "命中率：<2%",
            "2~4小时/场": "2~4小时/场",
            "4~6小时/场": "4~6小时/场",
            "6~8小时/场": "6~8小时/场",
            "2小时内/场": "2小时以内/场",
            "2小时以内/场": "2小时以内/场",
            "2小时以内": "2小时以内/场",
            "2H以内/场": "2小时以内/场",
            "2h以内/场": "2小时以内/场",
            "2H": "2小时以内/场",
            "2h": "2小时以内/场",
            "2~4H/场": "2~4小时/场",
            "2~4h/场": "2~4小时/场",
            "4~6H/场": "4~6小时/场",
            "4~6h/场": "4~6小时/场",
            "6~8H/场": "6~8小时/场",
            "6~8h/场": "6~8小时/场",
            "4H": "4~6小时/场",
            "4h": "4~6小时/场",
            "6H": "4~6小时/场",
            "6h": "4~6小时/场",
        }
        return _hour_map.get(v, v)

    print(f"  → 国内兼职: 行数={len(rows)}")
    await page.wait_for_timeout(500)

    # 按列标题定位（避免依赖 id 哈希）
    parttime_table = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("工作类型")')
    ).first
    if await parttime_table.count() == 0:
        print("    ⚠ 未找到兼职子表")
        return
    tbody = parttime_table.locator(".ant-table-tbody")

    for i, row_data in enumerate(rows):
        if i > 0:
            add_btn = (
                parttime_table.locator('[class*="addIcon"]').last
                if await parttime_table.locator('[class*="addIcon"]').count() > 0
                else page.locator('.anticon-plus-circle').last
            )
            await add_btn.click()
            await page.wait_for_timeout(500)

        table_rows = tbody.locator("tr.ant-table-row")
        row = table_rows.nth(i)

        col_map = {"工作类型": 0, "工作难度": 1}
        for field, col_idx in col_map.items():
            v = row_data.get(field)
            if v and str(v).strip():
                val = normalize_difficulty(str(v).strip()) if field == "工作难度" else str(v).strip()
                try:
                    selected_label = await ant_select_in_row(row, col_idx, val, page)
                    if field == "工作难度" and selected_label != val:
                        print(f"    ℹ 工作难度映射: {val} → {selected_label}")
                    if field == "工作类型":
                        await page.wait_for_timeout(1500)  # 等级联更新后才能选工作难度
                except Exception as e:
                    if field == "工作难度":
                        # 打印实际可见选项，帮助排查文本不一致
                        dd = page.locator('.ant-select-dropdown:not(.ant-select-dropdown-hidden)')
                        if await dd.count() > 0:
                            items = dd.locator('.ant-select-dropdown-menu-item')
                            opts = [await items.nth(j).inner_text() for j in range(await items.count())]
                            print(f"    ⚠ 工作难度实际选项: {opts}")
                        else:
                            print(f"    ⚠ 工作难度下拉未出现，请确认工作类型已正确选择")
                    print(f"    ⚠ {field} 填写异常: {e}")
                    await page.keyboard.press("Escape")

        for placeholders, field in [
            (("请输入场次", "请输入样本量"), "测试场次/样本量"),
            (("请输入底薪",), "底薪"),
            (("请输入提成",), "提成"),
            (("请输入总金额（元）",), "总金额(元)"),
        ]:
            v = row_data.get(field)
            if v is not None and str(v).strip():
                for placeholder in placeholders:
                    if await ant_number_in_row(row, placeholder, v, quiet=True):
                        break
                else:
                    await ant_number_in_row(row, placeholders[0], v)

        await page.wait_for_timeout(200)


async def fill_other(page: Page, data: dict):
    """填写【其他】子表（发包内容 + 数量 + 总金额）"""
    other = data.get("other", {})
    rows = other.get("rows", []) if isinstance(other, dict) else other
    if not rows:
        return

    print(f"  → 其他: 行数={len(rows)}")

    # 等其他子表展开（勾选后有动画）
    await page.wait_for_timeout(500)

    # 定位【其他】子表：包含"发包内容"列的表格
    other_table = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("发包内容")')
    ).first
    tc = await other_table.count()
    if tc == 0:
        print("    ⚠ 未找到【其他】子表")
        return

    tbody = other_table.locator('.ant-table-tbody')

    for i, row_data in enumerate(rows):
        if i > 0:
            # 优先在 other_table 内部找加号按钮
            add_btn = other_table.locator('[class*="addIcon"]').last
            if await add_btn.count() == 0:
                add_btn = other_table.locator('.anticon-plus-circle').last
            if await add_btn.count() == 0:
                add_btn = page.locator('.anticon-plus-circle').last
            await add_btn.click()
            await page.wait_for_timeout(400)

        table_rows = tbody.locator('tr.ant-table-row')
        row = table_rows.nth(i)
        print(f"    行{i+1}: {row_data}")

        # 发包内容（下拉，第0列）
        content_val = row_data.get("发包内容")
        if content_val:
            await ant_select_in_row(row, 0, content_val, page)

        # 数量：用 ant_number_in_row 统一处理（含模糊匹配）
        qty = row_data.get("数量")
        if qty is not None:
            await ant_number_in_row(row, "请输入样本量", qty)

        # 总金额
        total = row_data.get("总金额(元)")
        if total is not None:
            await ant_number_in_row(row, "请输入总金额（元）", total)

        await page.wait_for_timeout(200)


# ─────────────────────────────────────────────
# 第2页：收款人信息（批量上传）
# ─────────────────────────────────────────────

async def fill_page2(page: Page, payee_excel_path: str):
    """批量上传收款人明细 Excel"""
    print("→ 第2页：批量上传收款人信息...")

    if not Path(payee_excel_path).exists():
        print(f"  ⚠ 文件不存在: {payee_excel_path}")
        return

    # 点【+】批量上传按钮（aria 或 title）
    batch_btn = page.locator('.ant-upload').filter(
        has_text=""
    ).first

    # 直接找隐藏的 file input（批量上传按钮后面）
    # 通常点击触发 file input
    file_inputs = page.locator('input[type="file"]')
    count = await file_inputs.count()
    print(f"  找到 {count} 个 file input")

    # 批量上传通常是第二个 file input（第一个是第1页截图用的）
    upload_input = file_inputs.last
    await upload_input.set_input_files(payee_excel_path)
    await page.wait_for_timeout(2000)

    # 等待上传结果（可能有成功/失败弹窗）
    await page.wait_for_timeout(1000)
    print("  ✓ 批量上传完成")


# ─────────────────────────────────────────────
# 第3页：确认提交
# ─────────────────────────────────────────────

async def handle_page3(page: Page, split_note: dict):
    """处理第3页确认提交"""
    print("→ 第3页：确认提交...")

    # 检测是否出现「拆单备注」必填框
    split_area = await page.get_by_text("拆单备注").count()

    if split_area > 0:
        print("\n" + "="*55)
        print("⚠  检测到【拆单备注】必填项")
        print("   请直接在浏览器里填写拆单原因，然后手动点【提交】")
        print("   完成后回到这里按 Enter 关闭浏览器")
        print("="*55)
        # 截图供参考
        await page.screenshot(path="confirm_page3.png", full_page=True)
        print("  📸 第3页截图: confirm_page3.png")
        input("\n  （浏览器操作完成后）按 Enter 关闭浏览器...")
        return

    # 无拆单备注 → 正常确认提交流程
    await page.screenshot(path="confirm_page3.png", full_page=True)
    print(f"  📸 第3页截图: confirm_page3.png")
    print("  收款人和金额确认无误后输入 '提交'，或输入 '取消' 放弃")

    while True:
        cmd = input("  命令 [提交/取消]: ").strip()
        if cmd == "提交":
            submit_btn = page.locator('button').filter(has_text="提交").last
            await submit_btn.click()
            await page.wait_for_timeout(2000)
            print("  ✓ 已点击提交！")
            break
        elif cmd == "取消":
            print("  已取消，请手动操作")
            break
        else:
            print("  请输入 '提交' 或 '取消'")


# ─────────────────────────────────────────────
# 主流程
# ─────────────────────────────────────────────

async def run_single_order(data: dict):
    """执行单个报销单的全流程自动化（供 main.py 调用）"""
    from parser import build_expense_note as _build  # fallback
    import importlib

    from expense_note import build_standard_expense_note

    expense_note = build_standard_expense_note(data)

    # 构建与 fill_page1 兼容的 ov 结构
    ov_data = {
        "overview": {
            "project_id":           data.get("project_id"),
            "fanbao_series":        data.get("fanbao_series"),
            "fanbao_type":          data.get("fanbao_type"),
            "cost_center":          data.get("cost_center"),
            "screenshot_required":  data.get("screenshot_required", True),
            "screenshot_path":      data.get("screenshot_path"),
            "reimbursement_types":  data.get("reimbursement_types", []),
            "expense_note_manual":  expense_note,
        },
        "gift_common":   data.get("gift_common"),
        "gift_special":  data.get("gift_special") or [],
        "questionnaire": data.get("questionnaire") or [],
        "parttime":      data.get("parttime"),
        "other":         data.get("other"),
        "payee_rules":   data.get("payee_rules", {}),
        "split_note":    data.get("split_note", {}),
    }

    async with async_playwright() as pw:
        if CHROME_USER_DATA:
            browser = await pw.chromium.launch_persistent_context(
                user_data_dir=CHROME_USER_DATA,
                channel=BROWSER_CHANNEL or None,
                headless=False,
                slow_mo=SLOW_MO,
                args=["--start-maximized"],
            )
            page = browser.pages[0] if browser.pages else await browser.new_page()
        else:
            browser = await pw.chromium.launch(
                channel=BROWSER_CHANNEL or None,
                headless=False,
                slow_mo=SLOW_MO,
                args=["--start-maximized"],
            )
            context = await browser.new_context(viewport={"width": 1920, "height": 1080})
            page = await context.new_page()

        try:
            await login(page)
            # login() 已经导航到 EPC_URL 并等待表单加载，无需再次 goto

            # 重复报销检测
            duplicate = page.locator('[class*="theoryBox"]')
            if await duplicate.count() > 0:
                print("\n⚠ 警告：该项目已存在报销单（可能是拆单）！")
                txt = await duplicate.inner_text()
                print(f"  {txt[:300]}")
                cmd = input("  确认继续提报？[y/n]: ").strip().lower()
                if cmd != "y":
                    print("  已取消")
                    return

            await fill_page1(page, ov_data, expense_note)

            await page.screenshot(path=f"page1_{data.get('project_id','')}.png", full_page=True)
            print(f"\n📸 第1页截图: page1_{data.get('project_id','')}.png")
            print("  （如需上传产品确认截图，请先在浏览器中手动上传）")
            print("  请检查第1页内容，确认无误后，请你自己在浏览器里点击『下一步』进入第2页")
            cmd = input("完成后在此输入 y 继续（agent 将自动填写第2页）[y/停止]: ").strip().lower()
            if cmd != "y":
                input("按 Enter 退出...")
                return

            print("\n→ 等待你在浏览器中点击『下一步』进入第2页...")
            await page.wait_for_selector('th:has-text("真实姓名")', timeout=0)
            await page.wait_for_timeout(1000)
            print("→ 已检测到第2页，开始自动填写")

            # 第2页：按姓名匹配设置金额和类型
            await fill_page2_by_rules(page, data.get("payee_rules", {}))

            await page.screenshot(path=f"page2_{data.get('project_id','')}.png", full_page=True)
            print(f"📸 第2页截图: page2_{data.get('project_id','')}.png")
            print("  请检查第2页内容，确认无误后，请你自己在浏览器里点击『下一步』进入第3页")
            cmd = input("完成后在此输入 y 继续（agent 将自动处理第3页）[y/停止]: ").strip().lower()
            if cmd != "y":
                input("按 Enter 退出...")
                return

            print("\n→ 等待你在浏览器中点击『下一步』进入第3页...")
            await page.wait_for_selector('button:has-text("提交")', timeout=0)
            await page.wait_for_timeout(1000)
            print("→ 已检测到第3页，开始自动处理")

            await handle_page3(page, data.get("split_note", {}))

        finally:
            input("\n按 Enter 关闭浏览器...")
            await browser.close()


async def _delete_payee_row(page: Page, payee_wrapper, row_index: int, name: str) -> bool:
    """删除 fixed-right 操作列同索引行的删除按钮（红色垃圾桶）"""
    # 优先候选：Ant Design 图标类 + 常见操作列结构
    right_row = payee_wrapper.locator(
        '.ant-table-fixed-right .ant-table-tbody tr.ant-table-row'
    ).nth(row_index)

    candidates = [
        '[title="删除"]',
        '[aria-label="删除"]',
        '.anticon-delete',
        '.anticon-rest',
        'i[class*="delete"]',
        'i[class*="trash"]',
        'button[class*="danger"]',
        'a:has-text("删除")',
        'button:has-text("删除")',
    ]
    for sel in candidates:
        btn = right_row.locator(sel).first
        try:
            if await btn.count() > 0:
                await btn.scroll_into_view_if_needed()
                await btn.click()
                await page.wait_for_timeout(400)
                # 可能弹二次确认（"确定"/"是"）
                for ok_sel in (
                    '.ant-popover-buttons button.ant-btn-primary',
                    '.ant-modal-footer button.ant-btn-primary',
                    'button:has-text("确定")',
                    'button:has-text("确认")',
                ):
                    try:
                        ok = page.locator(ok_sel).last
                        if await ok.count() > 0 and await ok.is_visible():
                            await ok.click()
                            break
                    except Exception:
                        pass
                await page.wait_for_timeout(400)
                print(f"  🗑 已删除多余行: {name}")
                return True
        except Exception:
            continue

    # 兜底①：class 名识别不了（自定义红色图标），扫 background-color / svg fill / ::before 伪元素背景，找红色系元素坐标点击
    try:
        box = await right_row.evaluate("""(rowEl) => {
            function isRed(colorStr) {
                if (!colorStr) return false;
                const m = colorStr.match(/rgba?\\(([\\d.]+),\\s*([\\d.]+),\\s*([\\d.]+)(?:,\\s*([\\d.]+))?\\)/);
                if (!m) return false;
                const r = +m[1], g = +m[2], b = +m[3];
                const a = m[4] !== undefined ? +m[4] : 1;
                if (a < 0.15) return false;
                return r > 140 && (r - g) > 30 && (r - b) > 30;
            }
            const all = rowEl.querySelectorAll('*');
            let best = null, bestScore = -1;
            for (const el of all) {
                const rect = el.getBoundingClientRect();
                if (rect.width < 3 || rect.width > 70 || rect.height < 3 || rect.height > 70) continue;
                const cs = getComputedStyle(el);
                let score = 0;
                if (isRed(cs.backgroundColor)) score = Math.max(score, 3);
                if (isRed(cs.color)) score = Math.max(score, 2);
                try {
                    const before = getComputedStyle(el, '::before');
                    if (isRed(before.backgroundColor) || isRed(before.color)) score = Math.max(score, 2);
                } catch (e) {}
                if (el.tagName === 'path' || el.tagName === 'svg') {
                    const fill = el.getAttribute('fill');
                    if (isRed(fill)) score = Math.max(score, 3);
                }
                if (score > bestScore) {
                    bestScore = score;
                    best = {x: rect.left + rect.width / 2, y: rect.top + rect.height / 2};
                }
            }
            return bestScore > 0 ? best : null;
        }""")
        if box:
            await page.mouse.click(box["x"], box["y"])
            await page.wait_for_timeout(400)
            for ok_sel in (
                '.ant-popover-buttons button.ant-btn-primary',
                '.ant-modal-footer button.ant-btn-primary',
                'button:has-text("确定")',
                'button:has-text("确认")',
            ):
                try:
                    ok = page.locator(ok_sel).last
                    if await ok.count() > 0 and await ok.is_visible():
                        await ok.click()
                        break
                except Exception:
                    pass
            await page.wait_for_timeout(400)
            print(f"  🗑 已删除多余行(颜色识别): {name}")
            return True
    except Exception as e:
        print(f"  ⚠ 颜色识别删除按钮异常: {e}")

    # 兜底②：位置识别 —— 操作列固定顺序「编辑/删除/重置」共3个图标，取中间一个
    try:
        box2 = await right_row.evaluate("""(rowEl) => {
            const cells = rowEl.querySelectorAll('td');
            const opCell = cells[cells.length - 1];
            if (!opCell) return null;
            const cand = Array.from(opCell.querySelectorAll('svg, i, span, button, a'));
            const seen = new Set();
            const uniq = [];
            for (const el of cand) {
                const rect = el.getBoundingClientRect();
                if (rect.width < 4 || rect.width > 60 || rect.height < 4 || rect.height > 60) continue;
                const key = Math.round(rect.left) + ',' + Math.round(rect.top);
                if (seen.has(key)) continue;
                seen.add(key);
                uniq.push(rect);
            }
            if (uniq.length < 2) return null;
            const mid = uniq[Math.floor(uniq.length / 2)];
            return {x: mid.left + mid.width / 2, y: mid.top + mid.height / 2, count: uniq.length};
        }""")
        if box2:
            await page.mouse.click(box2["x"], box2["y"])
            await page.wait_for_timeout(400)
            for ok_sel in (
                '.ant-popover-buttons button.ant-btn-primary',
                '.ant-modal-footer button.ant-btn-primary',
                'button:has-text("确定")',
                'button:has-text("确认")',
            ):
                try:
                    ok = page.locator(ok_sel).last
                    if await ok.count() > 0 and await ok.is_visible():
                        await ok.click()
                        break
                except Exception:
                    pass
            await page.wait_for_timeout(400)
            print(f"  🗑 已删除多余行(位置识别 共{box2['count']}个图标): {name}")
            return True
    except Exception as e:
        print(f"  ⚠ 位置识别删除按钮异常: {e}")

    # 全部失败：打印该行操作列 HTML 供排查
    try:
        html = await right_row.locator("td").last.inner_html()
        print(f"  🔍 {name} 操作列HTML（供排查）: {html[:500]}")
    except Exception:
        pass

    return False


async def _payee_total_pages(payee_wrapper) -> int:
    """收款人表格自身的分页（页码 1,2,3...N），返回总页数"""
    pagination = payee_wrapper.locator('.ant-pagination').last
    if await pagination.count() == 0:
        return 1
    items = pagination.locator('.ant-pagination-item')
    cnt = await items.count()
    if cnt == 0:
        return 1
    try:
        last_txt = (await items.nth(cnt - 1).inner_text()).strip()
        return max(1, int(last_txt))
    except Exception:
        return 1


async def _goto_payee_page(payee_wrapper, page: Page, page_num: int):
    """跳转收款人表格自身分页到第 page_num 页"""
    pagination = payee_wrapper.locator('.ant-pagination').last
    if await pagination.count() == 0:
        return
    item = pagination.locator(f'.ant-pagination-item-{page_num}')
    if await item.count() == 0:
        item = pagination.get_by_text(str(page_num), exact=True).first
    if await item.count() > 0:
        await item.click()
        await page.wait_for_timeout(600)


async def _collect_payee_names_all_pages(payee_wrapper, page: Page, total_pages: int) -> list:
    """遍历收款人表格所有分页，采集姓名列表（顺序=页码顺序）"""
    rows = await _collect_payee_rows_all_pages(payee_wrapper, page, total_pages)
    return [r["name"] for r in rows]


async def _payee_phone_col(payee_wrapper) -> int:
    """定位「联系方式」列在主表(scroll body)表头里的序号，找不到返回 -1"""
    heads = payee_wrapper.locator('.ant-table-scroll thead th')
    cnt = await heads.count()
    for i in range(cnt):
        try:
            txt = (await heads.nth(i).inner_text()).strip()
        except Exception:
            txt = ""
        if "联系方式" in txt or "手机" in txt or "电话" in txt:
            return i
    return -1


async def _collect_payee_rows_all_pages(payee_wrapper, page: Page, total_pages: int) -> list:
    """遍历收款人表格所有分页，采集 [{name, phone}] 列表（顺序=页码顺序）
    姓名在 fixed-left（真实姓名列），手机号（联系方式）在可滚动的主表区，两边分别按行序号取值
    """
    phone_col = await _payee_phone_col(payee_wrapper)
    rows = []
    # 调用方可能刚停留在最后一页；必须先回到第一页再按页遍历，
    # 否则会把最后页当成第1页，产生重复/漏读。
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    for pg in range(1, total_pages + 1):
        if pg > 1:
            await _goto_payee_page(payee_wrapper, page, pg)
        fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
        main_rows = payee_wrapper.locator('.ant-table-scroll .ant-table-body .ant-table-tbody tr.ant-table-row')
        cnt = await fixed_rows.count()
        main_cnt = await main_rows.count()
        for i in range(cnt):
            try:
                nm = (await fixed_rows.nth(i).locator("td").nth(1).inner_text()).strip()
            except Exception:
                nm = ""
            ph = ""
            if phone_col >= 0 and i < main_cnt:
                try:
                    ph = (await main_rows.nth(i).locator("td").nth(phone_col).inner_text()).strip()
                except Exception:
                    ph = ""
            rows.append({"name": nm, "phone": ph})
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    return rows


def _normalize_person_name(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _person_name_candidates(value: str) -> list[str]:
    """拆分“黄蔚（梁阳）”或“黄蔚/梁阳”这类二选一姓名，保留原文兜底。"""
    raw = str(value or "").strip()
    if not raw:
        return []
    candidates = [raw]
    parenthesized = re.fullmatch(r"\s*([^（）()]+?)\s*[（(]\s*([^（）()]+?)\s*[）)]\s*", raw)
    if parenthesized:
        candidates.extend(parenthesized.groups())
    elif "/" in raw or "／" in raw:
        candidates.extend(part.strip() for part in re.split(r"[/／]", raw))
    result = []
    seen = set()
    for item in candidates:
        key = _normalize_person_name(item)
        if key and key not in seen:
            seen.add(key)
            result.append(item.strip())
    return result


def _manual_alias_records(payee_rules: dict) -> dict[str, dict]:
    records = {}
    for item in payee_rules.get("_manual_aliases") or []:
        source_key = _normalize_person_name(item.get("source_name"))
        platform_name = str(item.get("platform_name") or "").strip()
        if source_key and platform_name:
            records[source_key] = {
                "platform_name": platform_name,
                "phone": str(item.get("phone") or "").strip(),
            }
    return records


def _fuzzy_name_suggestion(source_name: str, platform_name: str) -> dict | None:
    """仅生成候选建议，不自动把相似姓名视为同一人。"""
    source = _normalize_person_name(source_name)
    platform = _normalize_person_name(platform_name)
    if not source or not platform or source == platform:
        return None
    source_pinyin = "".join(lazy_pinyin(source))
    platform_pinyin = "".join(lazy_pinyin(platform))
    if source_pinyin and source_pinyin == platform_pinyin:
        return {"score": 98, "reason": "姓名拼音一致（可能为同音字）"}
    ratio = SequenceMatcher(None, source, platform).ratio()
    same_surname = len(source) >= 2 and len(platform) >= 2 and source[0] == platform[0]
    same_chars = len(set(source) & set(platform))
    if same_surname and len(source) == len(platform) and same_chars >= len(source) - 1:
        return {"score": 82, "reason": "同姓且仅一字不同"}
    if same_surname and ratio >= 0.67:
        return {"score": round(ratio * 100), "reason": "同姓且姓名相似"}
    return None


def _proxy_value(item: dict, *keys: str) -> str:
    for key in keys:
        value = item.get(key)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _assignment_for_source(payee_rules: dict, name: str) -> dict | None:
    specific = {
        _normalize_person_name(item.get("name")): item
        for item in payee_rules.get("specific") or []
        if _normalize_person_name(item.get("name"))
    }
    key = _normalize_person_name(name)
    if key in specific:
        item = specific[key]
        return {
            "type": item.get("type") or "玩家",
            "amount": item.get("amount"),
            "source_name": name,
            "source_phone": (payee_rules.get("known_phones") or {}).get(name, ""),
        }
    if key in {_normalize_person_name(player) for player in payee_rules.get("known_players") or []}:
        return {
            "type": "玩家",
            "amount": payee_rules.get("default_player_amount"),
            "source_name": name,
            "source_phone": (payee_rules.get("known_phones") or {}).get(name, ""),
        }
    return None


def _proxy_records(payee_rules: dict) -> list[dict]:
    records = []
    for item in payee_rules.get("proxies") or []:
        source_name = _proxy_value(item, "source_name", "source", "payer_name", "from_name")
        proxy_name = _proxy_value(item, "proxy_name", "proxy", "collector_name", "to_name")
        if not source_name or not proxy_name or _normalize_person_name(source_name) == _normalize_person_name(proxy_name):
            continue
        records.append({
            "source_name": source_name,
            "source_phone": _proxy_value(item, "source_phone", "payer_phone", "from_phone"),
            "proxy_name": proxy_name,
            "proxy_phone": _proxy_value(item, "proxy_phone", "collector_phone", "to_phone"),
            "reason": _proxy_value(item, "reason", "relation") or "代收",
        })
    return records


async def _payee_header_texts(payee_wrapper) -> list[str]:
    headers = payee_wrapper.locator('.ant-table-scroll thead th')
    count = await headers.count()
    values = []
    for index in range(count):
        try:
            values.append((await headers.nth(index).inner_text()).strip())
        except Exception:
            values.append("")
    return values


def _header_index(headers: list[str], *alternatives: tuple[str, ...]) -> int:
    for index, header in enumerate(headers):
        compact = re.sub(r"\s+", "", header)
        for terms in alternatives:
            if all(term in compact for term in terms):
                return index
    return -1


async def _cell_text_or_input(cell) -> str:
    try:
        inp = cell.locator("input").first
        if await inp.count() > 0:
            value = await inp.evaluate("el => el.value")
            if value not in (None, ""):
                return str(value).strip()
        return (await cell.inner_text()).strip()
    except Exception:
        return ""


async def _row_validation_messages(main_row) -> list[str]:
    try:
        return await main_row.evaluate("""row => {
            const seen = new Set();
            const messages = [];
            for (const el of row.querySelectorAll('.ant-form-explain, .ant-form-item-explain-error, [class*="form-explain"], [class*="has-error"]')) {
                const text = (el.innerText || el.textContent || '').trim();
                if (text && !seen.has(text)) { seen.add(text); messages.push(text); }
            }
            for (const input of row.querySelectorAll('input')) {
                const parent = input.closest('.ant-form-item, .ant-input, .ant-input-number');
                const cls = ((input.className || '') + ' ' + (parent?.className || '')).toLowerCase();
                if (cls.includes('error') && !messages.includes('字段显示红色校验错误')) messages.push('字段显示红色校验错误');
            }
            return messages;
        }""")
    except Exception:
        return []


async def _fixed_name_validation_messages(fixed_row, name: str) -> list[str]:
    """识别 EPC 固定左侧“真实姓名”列中以红色显示的实名错误。"""
    try:
        red_name = await fixed_row.evaluate(r"""(row, expectedName) => {
            const normalize = value => String(value || '').replace(/\\s+/g, '').trim();
            const expected = normalize(expectedName);
            const isRed = color => {
                const match = String(color || '').match(/rgba?\((\d+),\s*(\d+),\s*(\d+)/i);
                if (!match) return false;
                const red = Number(match[1]), green = Number(match[2]), blue = Number(match[3]);
                return red >= 150 && green <= 130 && blue <= 130 && red - green >= 45;
            };
            for (const element of row.querySelectorAll('*')) {
                const text = normalize(element.innerText || element.textContent);
                if (text !== expected) continue;
                if (isRed(getComputedStyle(element).color)) return true;
            }
            return false;
        }""", name)
        return ["真实姓名显示红色，实名信息有误"] if red_name else []
    except Exception:
        return []


async def _collect_payee_platform_rows(payee_wrapper, page: Page, total_pages: int) -> list[dict]:
    """跨分页读取收款人行及平台字段完整性。"""
    headers = await _payee_header_texts(payee_wrapper)
    columns = {
        "phone": _header_index(headers, ("联系方式",), ("手机",), ("电话",)),
        "type": _header_index(headers, ("收款人", "类型"), ("类型",)),
        "amount": _header_index(headers, ("礼金", "金额")),
        "alipay_login": _header_index(headers, ("支付宝", "登录"), ("支付宝", "账号"), ("支付宝", "账户")),
        "alipay_name": _header_index(headers, ("支付宝", "实名"), ("支付宝", "姓名"), ("实名", "姓名")),
    }
    rows: list[dict] = []
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    for page_number in range(1, total_pages + 1):
        if page_number > 1:
            await _goto_payee_page(payee_wrapper, page, page_number)
        fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
        main_rows = payee_wrapper.locator('.ant-table-scroll .ant-table-body .ant-table-tbody tr.ant-table-row')
        count = await fixed_rows.count()
        main_count = await main_rows.count()
        for index in range(count):
            try:
                name = (await fixed_rows.nth(index).locator("td").nth(1).inner_text()).strip()
            except Exception:
                name = ""
            row = {
                "name": name,
                "phone": "",
                "type": "",
                "actual_amount": None,
                "alipay_login": None,
                "alipay_name": None,
                "validation_errors": [],
                "page": page_number,
            }
            if index < main_count:
                main_row = main_rows.nth(index)
                cells = main_row.locator("td")
                for field in ("phone", "type", "alipay_login", "alipay_name"):
                    column = columns[field]
                    if column >= 0:
                        row[field] = await _cell_text_or_input(cells.nth(column))
                if columns["amount"] >= 0:
                    row["actual_amount"] = _amount_value(await _cell_text_or_input(cells.nth(columns["amount"])))
                row["validation_errors"] = await _row_validation_messages(main_row)
            for message in await _fixed_name_validation_messages(fixed_rows.nth(index), name):
                if message not in row["validation_errors"]:
                    row["validation_errors"].append(message)
            rows.append(row)
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    return rows


def _row_matches(row: dict, name: str, phone: str = "") -> bool:
    if phone and row.get("phone") == phone:
        return True
    return _normalize_person_name(row.get("name")) == _normalize_person_name(name)


async def _apply_payee_assignment(page: Page, main_row, payee_type: str, amount, type_column: int, amount_column: int) -> bool:
    try:
        cells = main_row.locator("td")
        type_cell = cells.nth(type_column) if type_column >= 0 else main_row
        type_select = type_cell.locator('.ant-select-selection').first
        if await type_select.count() == 0:
            type_select = main_row.locator('.ant-select-selection').first
        await type_select.click()
        await page.wait_for_selector('.ant-select-dropdown:not(.ant-select-dropdown-hidden)', timeout=5_000)
        option = page.locator(
            f'.ant-select-dropdown:not(.ant-select-dropdown-hidden) .ant-select-dropdown-menu-item:has-text("{payee_type}")'
        ).first
        if await option.count() == 0:
            await page.keyboard.press("Escape")
            return False
        await option.click()
        await page.wait_for_timeout(250)
        target = cells.nth(amount_column) if amount_column >= 0 else main_row
        input_box = target.locator('input').first
        if await input_box.count() == 0:
            inputs = main_row.locator('input')
            input_count = await inputs.count()
            for index in range(input_count):
                candidate = inputs.nth(index)
                placeholder = (await candidate.get_attribute('placeholder') or "")
                if "金额" in placeholder:
                    input_box = candidate
                    break
        if await input_box.count() == 0:
            return False
        await input_box.evaluate("""(input, value) => {
            const setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
            setter.call(input, String(value));
            input.dispatchEvent(new Event('input', {bubbles:true}));
            input.dispatchEvent(new Event('change', {bubbles:true}));
            input.blur();
        }""", amount)
        return True
    except Exception:
        try:
            await page.keyboard.press("Escape")
        except Exception:
            pass
        return False


async def fill_page2_by_rules(page: Page, payee_rules: dict, runtime: Any = None) -> dict:
    """按草稿规则填写收款人，并处理代收、平台多余人员和字段完整性。"""
    print("→ 第2页：设置收款人金额...")
    await page.wait_for_selector(".ant-table-tbody tr", timeout=TIMEOUT)
    await page.wait_for_timeout(500)
    payee_wrapper = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("真实姓名")')
    ).first
    if await payee_wrapper.count() == 0:
        payee_wrapper = page.locator('.ant-table-wrapper').last

    total_pages = await _payee_total_pages(payee_wrapper)
    platform_rows = await _collect_payee_platform_rows(payee_wrapper, page, total_pages)
    print(f"  收款人表格共 {total_pages} 页，平台收款人 {len(platform_rows)} 人")

    known_phones = payee_rules.get("known_phones") or {}
    direct_names = list(payee_rules.get("known_players") or []) + [
        item.get("name", "") for item in payee_rules.get("specific") or []
    ]
    proxies = _proxy_records(payee_rules)
    proxy_sources = {_normalize_person_name(item["source_name"]) for item in proxies}
    assignments: dict[str, dict] = {}
    assignment_names: dict[str, str] = {}
    alias_to_assignment: dict[str, str] = {}
    manual_aliases = _manual_alias_records(payee_rules)
    expected_phones: dict[str, str] = {}
    proxy_issues: list[dict] = []

    for name in direct_names:
        key = _normalize_person_name(name)
        if not key or key in proxy_sources:
            continue
        assignment = _assignment_for_source(payee_rules, name)
        if assignment:
            assignments[key] = assignment
            assignment_names[key] = name
            expected_phones[key] = assignment.get("source_phone", "")
            for candidate in _person_name_candidates(name):
                alias_to_assignment.setdefault(_normalize_person_name(candidate), key)
    for source_key, alias in manual_aliases.items():
        if source_key in assignments:
            alias_to_assignment[_normalize_person_name(alias["platform_name"])] = source_key

    delete_names: list[str] = []
    for proxy in proxies:
        source_assignment = _assignment_for_source(payee_rules, proxy["source_name"])
        source_key = _normalize_person_name(proxy["source_name"])
        target_key = _normalize_person_name(proxy["proxy_name"])
        if not source_assignment:
            proxy_issues.append({"message": f"代收来源 {proxy['source_name']} 在草稿中没有可用金额规则"})
            continue
        if target_key in assignments:
            proxy_issues.append({"message": f"代收人 {proxy['proxy_name']} 同时有独立收款规则，无法自动合并金额"})
            continue
        assignments[target_key] = {**source_assignment, "source_name": proxy["source_name"], "proxy": True}
        expected_phones[target_key] = proxy.get("proxy_phone") or ""
        source_rows = [row for row in platform_rows if _row_matches(row, proxy["source_name"], proxy.get("source_phone", ""))]
        target_rows = [row for row in platform_rows if _row_matches(row, proxy["proxy_name"], proxy.get("proxy_phone", ""))]
        if source_rows and target_rows:
            delete_names.extend(row["name"] for row in source_rows if row.get("name"))
        elif source_rows and not target_rows:
            proxy_issues.append({"message": f"代收关系：平台出现 {proxy['source_name']}，但未出现代收人 {proxy['proxy_name']}；请补充或确认代收人"})
        elif not source_rows and not target_rows:
            proxy_issues.append({"message": f"代收关系：平台未出现 {proxy['source_name']} 或代收人 {proxy['proxy_name']}"})

    def canonical_key(row: dict) -> str:
        name_key = _normalize_person_name(row.get("name"))
        if name_key in alias_to_assignment:
            return alias_to_assignment[name_key]
        for key, phone in expected_phones.items():
            if phone and row.get("phone") == phone:
                return key
        return ""

    # 双姓名候选组只使用 EPC 列表中最先出现的那个名字；后续候选行仍作为多余人员交由人工确认。
    chosen_platform_names: dict[str, str] = {}
    for row in platform_rows:
        key = canonical_key(row)
        if key and key not in chosen_platform_names:
            chosen_platform_names[key] = _normalize_person_name(row.get("name"))

    platform_keys = set(chosen_platform_names)
    missing = [key for key in assignments if key not in platform_keys]
    extras = []
    for row in platform_rows:
        key = canonical_key(row)
        name_key = _normalize_person_name(row.get("name"))
        if key and chosen_platform_names.get(key) != name_key:
            extras.append({
                **row,
                "review_reason": f"双姓名候选：已按 EPC 顺序选择 {assignment_names.get(key, key)} 的首个匹配人员",
            })
        elif not key and name_key not in proxy_sources:
            extras.append(row)
    review_rows = [{
        "name": row.get("name", ""),
        "phone": row.get("phone", ""),
        "action": "delete",
        "reason": row.get("review_reason") or "平台有但草稿无",
    } for row in extras]
    possible_matches = []
    for missing_key in missing:
        source_name = assignment_names.get(missing_key, missing_key)
        for row in extras:
            suggestions = []
            for candidate_name in _person_name_candidates(source_name):
                suggestion = _fuzzy_name_suggestion(candidate_name, row.get("name", ""))
                if suggestion:
                    suggestions.append((suggestion, candidate_name))
            if suggestions:
                suggestion, matched_candidate = max(suggestions, key=lambda item: item[0].get("score", 0))
                possible_matches.append({
                    "source_name": source_name,
                    "platform_name": row.get("name", ""),
                    "phone": row.get("phone", ""),
                    "matched_candidate": matched_candidate,
                    **suggestion,
                })

    keep_rules: dict[str, dict] = {}
    if review_rows or proxy_issues or missing:
        print("__PAYEE_REVIEW__" + json.dumps({
            "platform_extra": review_rows,
            "missing": [{"name": assignment_names.get(name, name), "reason": "草稿有但平台无"} for name in missing],
            "proxy_issues": proxy_issues,
            "possible_matches": possible_matches,
        }, ensure_ascii=False))
        raw = (await _ask_runtime(
            runtime,
            "请核对平台多余人员、草稿缺失人员及代收关系：默认删除明确多余人员；保留时必须填写身份和金额。确认后继续填写已匹配人员。",
            "REQUIRES_ATTENTION",
        )).strip()
        if raw.lower() == "n":
            return {"stopped": True, "missing": missing, "extras": review_rows, "proxy_issues": proxy_issues}
        decisions = {}
        alias_associations = []
        if raw.lower() == "y":
            decisions = {row["name"]: {"action": "delete"} for row in review_rows}
        else:
            try:
                payload = json.loads(raw)
                decisions = {item.get("name"): item for item in payload.get("platform_extra", []) if item.get("name")}
                alias_associations = payload.get("alias_associations") or []
            except Exception as exc:
                print(f"  ⚠ 无法解析平台审核结果: {exc}")
                return {"stopped": True, "missing": missing, "extras": review_rows, "proxy_issues": proxy_issues}
        seen_alias_platforms = set()
        for association in alias_associations:
            source_key = _normalize_person_name(association.get("source_name"))
            platform_name = str(association.get("platform_name") or "").strip()
            platform_key = _normalize_person_name(platform_name)
            if not source_key or not platform_key or source_key not in assignments or source_key not in missing:
                proxy_issues.append({"message": f"无效的同人关联：{association.get('source_name') or '未命名'} → {platform_name or '未命名'}"})
                continue
            if platform_key in seen_alias_platforms:
                proxy_issues.append({"message": f"同一 EPC 人员不能同时关联多个 JSON 人员：{platform_name}"})
                continue
            seen_alias_platforms.add(platform_key)
            alias_to_assignment[platform_key] = source_key
            chosen_platform_names[source_key] = platform_key
            missing.remove(source_key)
            decisions[platform_name] = {"action": "associate", "target_name": assignment_names.get(source_key, source_key)}
            manual_aliases_list = payee_rules.setdefault("_manual_aliases", [])
            manual_aliases_list[:] = [
                item for item in manual_aliases_list
                if _normalize_person_name(item.get("source_name")) != source_key
            ]
            manual_aliases_list.append({
                "source_name": assignment_names.get(source_key, source_key),
                "platform_name": platform_name,
                "phone": association.get("phone") or "",
                "reason": association.get("reason") or "用户确认疑似同一人",
            })
            if runtime is not None:
                runtime.log(f"✓ 已确认同人关联：{assignment_names.get(source_key, source_key)} → {platform_name}")
                runtime.store.update_order(runtime.order_id, draft=runtime.data)
        for row in review_rows:
            decision = decisions.get(row["name"], {"action": "delete"})
            if decision.get("action") == "associate":
                continue
            if decision.get("action", "delete") == "keep":
                payee_type = decision.get("type") or ""
                amount = _amount_value(decision.get("amount"))
                if payee_type not in {"玩家", "渠道接口人", "兼职"} or amount is None or amount <= 0:
                    proxy_issues.append({"message": f"保留人员 {row['name']} 缺少有效身份或金额"})
                else:
                    keep_rules[_normalize_person_name(row["name"])] = {"type": payee_type, "amount": amount, "source_name": row["name"]}
                    # 人工保留优先于 JSON：写入当前单的运行规则，后续自动核验/回滚不得删除它。
                    manual_kept = payee_rules.setdefault("_manual_kept", [])
                    manual_kept[:] = [
                        item for item in manual_kept
                        if _normalize_person_name(item.get("name")) != _normalize_person_name(row["name"])
                    ]
                    manual_kept.append({
                        "name": row["name"],
                        "phone": row.get("phone", ""),
                        "type": payee_type,
                        "amount": amount,
                        "reason": "用户手动保留，优先于 JSON 自动回滚",
                    })
                    if runtime is not None:
                        try:
                            # 将人工决定保存到当前单草稿，服务重启/刷新后仍保持优先级。
                            runtime.store.update_order(runtime.order_id, draft=runtime.data)
                            runtime.log(f"✓ 已记录人工保留：{row['name']}（{payee_type} {amount:g}元），优先于 JSON")
                        except Exception as exc:
                            runtime.log(f"⚠ 人工保留未能持久化到草稿：{exc}")
            else:
                delete_names.append(row["name"])
        if proxy_issues:
            return {"stopped": True, "missing": missing, "extras": review_rows, "proxy_issues": proxy_issues}
        if missing and runtime is not None:
            runtime.log("⚠ 已由人工确认草稿人员未在 EPC 匹配，继续填写其余人员；回读核验时仍会提示未匹配人员。")

    for target_name in list(dict.fromkeys(name for name in delete_names if name)):
        found = False
        total_pages = await _payee_total_pages(payee_wrapper)
        for page_number in range(1, total_pages + 1):
            await _goto_payee_page(payee_wrapper, page, page_number)
            fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
            count = await fixed_rows.count()
            for index in range(count):
                try:
                    actual_name = (await fixed_rows.nth(index).locator("td").nth(1).inner_text()).strip()
                except Exception:
                    actual_name = ""
                if _normalize_person_name(actual_name) == _normalize_person_name(target_name):
                    if not await _delete_payee_row(page, payee_wrapper, index, actual_name):
                        return {"stopped": True, "missing": missing, "extras": review_rows, "proxy_issues": [{"message": f"无法自动删除 {target_name}"}]}
                    await page.wait_for_timeout(500)
                    found = True
                    break
            if found:
                break
        if not found:
            print(f"  ⚠ 未找到待删除人员 {target_name}")

    total_pages = await _payee_total_pages(payee_wrapper)
    platform_rows = await _collect_payee_platform_rows(payee_wrapper, page, total_pages)
    headers = await _payee_header_texts(payee_wrapper)
    type_col = _header_index(headers, ("收款人", "类型"), ("类型",))
    amount_col = _header_index(headers, ("礼金", "金额"))
    matched_count = default_count = 0
    if runtime is not None:
        runtime.set_status("FILLING_PAYEES", "名单核对已确认，填写已匹配收款人的身份和金额")

    for page_number in range(1, total_pages + 1):
        await _checkpoint(runtime, f"收款人第 {page_number} 页填写前")
        await _goto_payee_page(payee_wrapper, page, page_number)
        fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
        main_rows = payee_wrapper.locator('.ant-table-scroll .ant-table-body .ant-table-tbody tr.ant-table-row')
        count = await fixed_rows.count()
        main_count = await main_rows.count()
        for index in range(count):
            if index >= main_count:
                continue
            try:
                name = (await fixed_rows.nth(index).locator("td").nth(1).inner_text()).strip()
            except Exception:
                name = ""
            matched = next((row for row in platform_rows if row.get("page") == page_number and _normalize_person_name(row.get("name")) == _normalize_person_name(name)), None)
            key = canonical_key(matched or {"name": name, "phone": ""})
            if key and chosen_platform_names.get(key) != _normalize_person_name(name):
                continue
            assignment = assignments.get(key) or keep_rules.get(_normalize_person_name(name))
            if not assignment:
                continue
            ok = await _apply_payee_assignment(page, main_rows.nth(index), assignment["type"], assignment["amount"], type_col, amount_col)
            if ok:
                if assignment.get("source_name") == name:
                    default_count += 1
                else:
                    matched_count += 1
                print(f"  第{page_number}页 {name} → {assignment['type']} {assignment['amount']}元 ✓")
            else:
                print(f"  ⚠ 无法填写 {name} 的类型或礼金金额")
            await _checkpoint(runtime, f"收款人第 {page_number} 页第 {index + 1} 行填写后")

    actual_rows = await _collect_payee_platform_rows(payee_wrapper, page, await _payee_total_pages(payee_wrapper))
    field_issues = []
    for row in actual_rows:
        if row.get("validation_errors"):
            field_issues.append({"name": row.get("name"), "field": "页面红色校验", "message": "；".join(row["validation_errors"]), "page": row.get("page")})
        if "玩家" in (row.get("type") or ""):
            if row.get("alipay_login") == "":
                field_issues.append({"name": row.get("name"), "field": "支付宝登录号", "message": "为空", "page": row.get("page")})
            if row.get("alipay_name") == "":
                field_issues.append({"name": row.get("name"), "field": "支付宝实名姓名", "message": "为空", "page": row.get("page")})
    if field_issues:
        # 不终止当前单；统一交给后续核验差异表在前端修正并直接重新核验。
        print("  ⚠ 已发现平台字段问题，将进入结构化核验修正流程")

    print(f"\n  ✓ 第2页处理完成（指定/代收/保留规则 {matched_count} 人，默认玩家 {default_count} 人）")
    return {"stopped": False, "missing": [], "extras": review_rows, "proxy_issues": [], "field_issues": []}
async def _header_column(payee_wrapper, keywords: tuple[str, ...]) -> int:
    headers = payee_wrapper.locator('.ant-table-scroll thead th')
    count = await headers.count()
    for index in range(count):
        try:
            text = (await headers.nth(index).inner_text()).strip()
        except Exception:
            text = ""
        if all(keyword in text for keyword in keywords):
            return index
    return -1


def _amount_value(value: Any) -> float | None:
    match = re.search(r"-?\d+(?:\.\d+)?", str(value or "").replace(",", ""))
    return float(match.group(0)) if match else None


async def read_actual_payees(page: Page) -> list[dict]:
    """跨所有分页回读 EPC 已填写的收款人类型和礼金金额。"""
    payee_wrapper = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("真实姓名")')
    ).first
    if await payee_wrapper.count() == 0:
        payee_wrapper = page.locator('.ant-table-wrapper').last

    total_pages = await _payee_total_pages(payee_wrapper)
    phone_col = await _payee_phone_col(payee_wrapper)
    type_col = await _header_column(payee_wrapper, ("类型",))
    amount_col = await _header_column(payee_wrapper, ("礼金", "金额"))
    rows: list[dict] = []
    # fill_page2_by_rules 完成后页面通常停在最后一页；核验必须从第一页开始。
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    for page_number in range(1, total_pages + 1):
        if page_number > 1:
            await _goto_payee_page(payee_wrapper, page, page_number)
        fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
        main_rows = payee_wrapper.locator('.ant-table-scroll .ant-table-body .ant-table-tbody tr.ant-table-row')
        count = await fixed_rows.count()
        main_count = await main_rows.count()
        for index in range(count):
            try:
                name = (await fixed_rows.nth(index).locator("td").nth(1).inner_text()).strip()
            except Exception:
                name = ""
            phone = ""
            payee_type = ""
            amount = None
            if index < main_count:
                main = main_rows.nth(index)
                cells = main.locator("td")
                if phone_col >= 0:
                    try:
                        phone = (await cells.nth(phone_col).inner_text()).strip()
                    except Exception:
                        pass
                if type_col >= 0:
                    try:
                        type_cell = cells.nth(type_col)
                        selected = type_cell.locator('.ant-select-selection-selected-value').first
                        payee_type = (await selected.inner_text()).strip() if await selected.count() else (await type_cell.inner_text()).strip()
                    except Exception:
                        pass
                if amount_col >= 0:
                    try:
                        amount_cell = cells.nth(amount_col)
                        input_box = amount_cell.locator("input").first
                        raw_amount = await input_box.evaluate("el => el.value") if await input_box.count() else await amount_cell.inner_text()
                        amount = _amount_value(raw_amount)
                    except Exception:
                        pass
            rows.append({
                "name": name,
                "phone": phone,
                "type": payee_type,
                "actual_amount": amount,
                "page": page_number,
            })
    if total_pages > 1:
        await _goto_payee_page(payee_wrapper, page, 1)
    return rows


def verify_actual_payees(
    payee_rules: dict,
    actual_rows: list[dict],
    expected_grand_total: float | int | None = None,
) -> list[dict]:
    """对比标准收款人规则与 EPC 实填结果，返回逐人核验明细。"""
    expected: list[dict] = []
    default_amount = payee_rules.get("default_player_amount")
    known_phones = payee_rules.get("known_phones") or {}
    for name in payee_rules.get("known_players") or []:
        expected.append({
            "payee_id": f"player:{name}",
            "name": name,
            "phone": known_phones.get(name, ""),
            "type": "玩家",
            "expected_amount": default_amount,
        })
    for item in payee_rules.get("specific") or []:
        expected.append({
            "payee_id": f"specific:{item.get('name', '')}:{item.get('type', '')}",
            "name": item.get("name", ""),
            "phone": item.get("phone", ""),
            "type": item.get("type", "玩家"),
            "expected_amount": item.get("amount"),
        })

    unused = set(range(len(actual_rows)))
    results: list[dict] = []
    for item in expected:
        candidates = []
        if item["phone"]:
            candidates = [index for index in unused if actual_rows[index].get("phone") == item["phone"]]
        if not candidates and item["name"]:
            candidates = [index for index in unused if actual_rows[index].get("name") == item["name"]]
        if not candidates:
            results.append({**item, "actual_amount": None, "result": "missing_in_epc"})
            continue
        index = candidates[0]
        unused.remove(index)
        actual = actual_rows[index]
        expected_amount = _amount_value(item.get("expected_amount"))
        actual_amount = actual.get("actual_amount")
        if len(candidates) > 1:
            result = "duplicate_in_epc"
        elif item.get("type") and item["type"] not in (actual.get("type") or ""):
            result = "type_mismatch"
        elif expected_amount is None or actual_amount is None or abs(expected_amount - actual_amount) > 0.01:
            result = "amount_mismatch"
        else:
            result = "matched"
        results.append({
            **item,
            "actual_amount": actual_amount,
            "actual_name": actual.get("name", ""),
            "actual_phone": actual.get("phone", ""),
            "actual_type": actual.get("type", ""),
            "page": actual.get("page"),
            "result": result,
        })
    for index in sorted(unused):
        actual = actual_rows[index]
        results.append({
            "payee_id": "",
            "name": actual.get("name", ""),
            "phone": actual.get("phone", ""),
            "type": actual.get("type", ""),
            "expected_amount": None,
            "actual_amount": actual.get("actual_amount"),
            "page": actual.get("page"),
            "result": "extra_in_epc",
        })
    if expected_grand_total is not None:
        expected_total = _amount_value(expected_grand_total)
        actual_total = sum(_amount_value(row.get("actual_amount")) or 0 for row in actual_rows)
        if expected_total is not None:
            results.append({
                "payee_id": "__total__",
                "name": "收款人实际合计",
                "phone": "",
                "type": "总金额对账",
                "expected_amount": expected_total,
                "actual_amount": actual_total,
                "page": 0,
                "result": "matched" if abs(expected_total - actual_total) <= 0.01 else "total_amount_mismatch",
            })
    return results


def verification_issue_messages(items: list[dict]) -> list[str]:
    """将核验结果转为前端执行日志可读的异常说明。"""
    labels = {
        "missing_in_epc": "草稿有但 EPC 未找到该收款人",
        "extra_in_epc": "EPC 有但草稿没有该收款人",
        "duplicate_in_epc": "EPC 存在重复匹配收款人",
        "type_mismatch": "收款人身份类型不一致",
        "amount_mismatch": "礼金金额不一致",
        "phone_mismatch": "手机号不一致",
        "alipay_login_missing": "支付宝登录号为空",
        "alipay_real_name_missing": "支付宝实名姓名为空",
        "platform_validation_error": "EPC 页面存在红色校验错误",
    }
    messages = []
    for item in items:
        result = item.get("result")
        if result == "matched":
            continue
        if result == "total_amount_mismatch":
            expected = _amount_value(item.get("expected_amount")) or 0
            actual = _amount_value(item.get("actual_amount")) or 0
            delta = actual - expected
            direction = "多" if delta > 0 else "少"
            messages.append(
                f"总金额异常：费用明细应报 {expected:g} 元，EPC 收款人实际合计 {actual:g} 元，{direction} {abs(delta):g} 元"
            )
            continue
        name = item.get("name") or item.get("actual_name") or "未命名人员"
        detail = labels.get(result, result or "未知异常")
        if result == "amount_mismatch":
            detail += f"（应填 {item.get('expected_amount')} 元，实际 {item.get('actual_amount')} 元）"
        elif result == "phone_mismatch":
            detail += f"（草稿 {item.get('phone') or '空'}，EPC {item.get('actual_phone') or '空'}）"
        elif result == "platform_validation_error" and item.get("platform_errors"):
            errors = item["platform_errors"]
            if any("真实姓名显示红色" in message for message in errors):
                detail = "实名信息有误（EPC 红名）"
            detail += f"（{'；'.join(errors)}）"
        messages.append(f"{name}：{detail}")
    return messages


def verification_review_payload(items: list[dict]) -> dict:
    """给前端渲染可编辑核验差异表的数据。"""
    review_items = []
    for item in items:
        review_items.append({
            "issue_key": _verification_issue_key(item),
            "payee_id": item.get("payee_id", ""),
            "name": item.get("name", ""),
            "phone": item.get("phone", ""),
            "actual_name": item.get("actual_name", item.get("name", "")),
            "actual_phone": item.get("actual_phone", item.get("phone", "")),
            "type": item.get("type", ""),
            "actual_type": item.get("actual_type", ""),
            "expected_amount": item.get("expected_amount"),
            "actual_amount": item.get("actual_amount"),
            "expected_alipay_login": item.get("expected_alipay_login", ""),
            "actual_alipay_login": item.get("actual_alipay_login", ""),
            "expected_alipay_name": item.get("expected_alipay_name", ""),
            "actual_alipay_name": item.get("actual_alipay_name", ""),
            "platform_errors": item.get("platform_errors", []),
            "page": item.get("page", 0),
            "result": item.get("result", ""),
        })
    return {"verification_issues": review_items}


def _verification_issue_key(item: dict) -> str:
    """生成稳定的核验问题标识，用于记录用户逐条忽略的项目。"""
    return "|".join([
        str(item.get("payee_id") or ""),
        str(item.get("result") or ""),
        _normalize_person_name(item.get("actual_name") or item.get("name") or ""),
        str(item.get("page") or ""),
    ])


async def _apply_text_cell(main_row, column: int, value: str, field_keywords: tuple[str, ...] = ()) -> bool:
    if value in (None, ""):
        return True
    try:
        input_box = main_row.locator("td").nth(column).locator("input").first if column >= 0 else main_row.locator("input").first
        if await input_box.count() == 0:
            inputs = main_row.locator("input")
            for index in range(await inputs.count()):
                candidate = inputs.nth(index)
                descriptor = " ".join(filter(None, [
                    await candidate.get_attribute("placeholder"),
                    await candidate.get_attribute("aria-label"),
                    await candidate.get_attribute("name"),
                ]))
                if field_keywords and all(keyword in descriptor for keyword in field_keywords):
                    input_box = candidate
                    break
        if await input_box.count() == 0:
            return False
        written = await input_box.evaluate("""(input, value) => {
            if (input.readOnly || input.disabled) return false;
            const setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
            setter.call(input, String(value));
            input.dispatchEvent(new Event('input', {bubbles:true}));
            input.dispatchEvent(new Event('change', {bubbles:true}));
            input.blur();
            return true;
        }""", value)
        return bool(written)
    except Exception:
        return False


async def _wait_for_payee_editor_closed(page: Page, timeout_ms: int = 5_000) -> bool:
    """等待编辑收款人的 Drawer/Modal 完全关闭，防止遮挡分页或后续点击。"""
    deadline = time.monotonic() + timeout_ms / 1000
    while time.monotonic() < deadline:
        try:
            editors = page.locator('.ant-drawer:visible, .ant-modal:visible')
            if await editors.count() == 0:
                return True
        except Exception:
            return False
        await page.wait_for_timeout(120)
    return False


async def _apply_payee_edit_dialog(
    page: Page,
    payee_wrapper,
    row_index: int,
    payee_name: str,
    alipay_login: str | None,
    alipay_name: str | None,
) -> list[str]:
    """通过 EPC 右侧“编辑收款人”弹窗填写支付宝字段。"""
    if not alipay_login and not alipay_name:
        return []
    errors = []
    right_row = payee_wrapper.locator(
        '.ant-table-fixed-right .ant-table-tbody tr.ant-table-row'
    ).nth(row_index)
    edit_button = None
    for selector in (
        '[title="编辑"]', '[aria-label="编辑"]', '.anticon-edit',
        'i[class*="edit"]', 'button:has-text("编辑")', 'a:has-text("编辑")',
    ):
        candidate = right_row.locator(selector).first
        if await candidate.count() > 0:
            edit_button = candidate
            break
    try:
        if edit_button is not None:
            await edit_button.click()
        else:
            edit_point = await right_row.evaluate("""(rowEl) => {
                const opCell = rowEl.querySelectorAll('td')[rowEl.querySelectorAll('td').length - 1];
                if (!opCell) return null;
                const nodes = Array.from(opCell.querySelectorAll('svg, i, span, button, a'))
                    .map(el => ({el, rect: el.getBoundingClientRect()}))
                    .filter(item => item.rect.width >= 4 && item.rect.width <= 70 && item.rect.height >= 4 && item.rect.height <= 70)
                    .sort((a, b) => a.rect.left - b.rect.left);
                if (!nodes.length) return null;
                const rect = nodes[0].rect;
                return {x: rect.left + rect.width / 2, y: rect.top + rect.height / 2};
            }""")
            if not edit_point:
                return [f"{payee_name}：未找到 EPC 的编辑按钮"]
            await page.mouse.click(edit_point["x"], edit_point["y"])

        # EPC 的“编辑收款人”在不同版本中可能是 Modal、Drawer 或普通浮层。
        # 不依赖容器 class，直接等待截图中实际存在的支付宝输入框。
        await page.wait_for_timeout(250)

        async def fill_field(label: str, value: str | None) -> Locator | None:
            if not value:
                return None
            field = page.locator(f'input[placeholder*="{label}"]:visible').last
            try:
                await field.wait_for(state="visible", timeout=5_000)
            except Exception:
                return None
            try:
                # 必须先选中输入框，再全选覆盖旧值，兼容 EPC 对焦点的要求。
                await field.click()
                await field.press("Control+A")
                await field.fill(str(value))
                await field.press("Tab")
                await field.evaluate("""input => {
                    input.dispatchEvent(new Event('input', {bubbles:true}));
                    input.dispatchEvent(new Event('change', {bubbles:true}));
                }""")
                return field
            except Exception:
                return None

        login_field = await fill_field("支付宝登录号", alipay_login)
        real_name_field = await fill_field("支付宝实名", alipay_name)
        if alipay_login and login_field is None:
            errors.append(f"{payee_name}：编辑弹窗中未找到支付宝登录号输入框")
        if alipay_name and real_name_field is None:
            errors.append(f"{payee_name}：编辑弹窗中未找到支付宝实名输入框")
        if errors:
            close = page.locator('button:has-text("取消"):visible').last
            if await close.count() > 0:
                await close.click()
            return errors

        # 优先在输入框所属编辑容器内找“保存/确认”，没有固定容器时再找全局可见按钮。
        anchor = login_field if login_field is not None else real_name_field
        editor = anchor.locator(
            "xpath=ancestor::*[contains(@class,'ant-drawer') or contains(@class,'ant-modal') or contains(@class,'ant-popover')][1]"
        ) if anchor is not None else page.locator("body")
        save = editor.locator('button:has-text("保存"):visible, button:has-text("确认"):visible, button:has-text("确定"):visible').last
        if await save.count() == 0:
            save = page.locator('button:has-text("保存"):visible, button:has-text("确认"):visible, button:has-text("确定"):visible').last
        if await save.count() == 0:
            return [f"{payee_name}：编辑弹窗中未找到保存按钮"]
        await save.click()
        if not await _wait_for_payee_editor_closed(page, timeout_ms=5_000):
            return [f"{payee_name}：已点击保存，但编辑面板未关闭；请在 EPC 确认保存后点击“已修正，重新核验”"]
        return []
    except Exception as exc:
        try:
            await page.keyboard.press("Escape")
        except Exception:
            pass
        return [f"{payee_name}：通过编辑弹窗填写支付宝信息失败：{exc}"]


async def apply_verification_corrections(page: Page, corrections: list[dict]) -> list[str]:
    """按前端差异表的修正值回写 EPC 当前收款人页。"""
    payee_wrapper = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("真实姓名")')
    ).first
    if await payee_wrapper.count() == 0:
        payee_wrapper = page.locator('.ant-table-wrapper').last
    headers = await _payee_header_texts(payee_wrapper)
    type_col = _header_index(headers, ("收款人", "类型"), ("类型",))
    amount_col = _header_index(headers, ("礼金", "金额"))
    alipay_login_col = _header_index(headers, ("支付宝", "登录"), ("支付宝", "账号"), ("支付宝", "账户"))
    alipay_name_col = _header_index(headers, ("支付宝", "实名"), ("支付宝", "姓名"), ("实名", "姓名"))
    errors = []
    for correction in corrections:
        if correction.get("result") == "total_amount_mismatch":
            continue
        target_name = correction.get("actual_name") or correction.get("name") or ""
        target_phone = correction.get("actual_phone") or correction.get("phone") or ""
        found = False
        total_pages = await _payee_total_pages(payee_wrapper)
        for page_number in range(1, total_pages + 1):
            await _goto_payee_page(payee_wrapper, page, page_number)
            fixed_rows = payee_wrapper.locator('.ant-table-fixed-left .ant-table-tbody tr.ant-table-row')
            main_rows = payee_wrapper.locator('.ant-table-scroll .ant-table-body .ant-table-tbody tr.ant-table-row')
            count = await fixed_rows.count()
            for index in range(count):
                try:
                    name = (await fixed_rows.nth(index).locator("td").nth(1).inner_text()).strip()
                except Exception:
                    name = ""
                if not _row_matches({"name": name, "phone": ""}, target_name, ""):
                    continue
                if correction.get("action") == "delete":
                    if not await _delete_payee_row(page, payee_wrapper, index, name):
                        errors.append(f"{target_name}：无法自动删除")
                    found = True
                    break
                if index >= await main_rows.count():
                    errors.append(f"{target_name}：未找到可编辑收款人行")
                    found = True
                    break
                main_row = main_rows.nth(index)
                type_ok = await _apply_payee_assignment(
                    page,
                    main_row,
                    correction.get("type") or correction.get("actual_type") or "玩家",
                    correction.get("amount"),
                    type_col,
                    amount_col,
                )
                if not type_ok:
                    errors.append(f"{target_name}：类型或礼金金额无法自动写入")
                if correction.get("alipay_login") or correction.get("alipay_name"):
                    errors.append(
                        f"{target_name}：支付宝登录号/实名请在 EPC 右侧“编辑”中手动保存，"
                        "完成后点击“已修正，重新核验”"
                    )
                found = True
                break
            if found:
                break
        if not found:
            errors.append(f"{target_name or '未命名人员'}：EPC 当前页未找到，无法应用修正")
    return errors


def automatic_verification_corrections(items: list[dict]) -> list[dict]:
    """可由 JSON 唯一确定的差异，优先自动回滚到草稿值。"""
    corrections = []
    for item in items:
        result = item.get("result")
        # 明确人工保留的人员不再受 JSON 自动回滚影响。
        if str(item.get("payee_id", "")).startswith("manual_keep:"):
            continue
        if result == "extra_in_epc":
            corrections.append({
                "action": "delete",
                "name": item.get("name", ""),
                "actual_name": item.get("actual_name") or item.get("name", ""),
                "actual_phone": item.get("actual_phone") or item.get("phone", ""),
                "result": result,
            })
        elif result in {"type_mismatch", "amount_mismatch", "platform_validation_error"}:
            # 草稿已给出身份与金额时，直接以草稿覆盖 EPC 当前行。
            if item.get("type") and item.get("expected_amount") is not None:
                corrections.append({
                    "action": "apply",
                    "payee_id": item.get("payee_id", ""),
                    "name": item.get("name", ""),
                    "phone": item.get("phone", ""),
                    "actual_name": item.get("actual_name") or item.get("name", ""),
                    "actual_phone": item.get("actual_phone") or item.get("phone", ""),
                    "type": item.get("type"),
                    "amount": item.get("expected_amount"),
                    "result": result,
                })
    return corrections


async def read_actual_payees(page: Page) -> list[dict]:
    """新版回读：统一读取手机号、类型、礼金、支付宝字段和红色校验。"""
    payee_wrapper = page.locator('.ant-table-wrapper').filter(
        has=page.locator('th:has-text("真实姓名")')
    ).first
    if await payee_wrapper.count() == 0:
        payee_wrapper = page.locator('.ant-table-wrapper').last
    return await _collect_payee_platform_rows(
        payee_wrapper,
        page,
        await _payee_total_pages(payee_wrapper),
    )


def _expected_payees_from_rules(payee_rules: dict) -> list[dict]:
    names = list(payee_rules.get("known_players") or []) + [
        item.get("name", "") for item in payee_rules.get("specific") or []
    ]
    proxy_by_source = {
        _normalize_person_name(item["source_name"]): item
        for item in _proxy_records(payee_rules)
    }
    manual_aliases = _manual_alias_records(payee_rules)
    expected, seen = [], set()
    for source_name in names:
        source_key = _normalize_person_name(source_name)
        if not source_key or source_key in seen:
            continue
        seen.add(source_key)
        assignment = _assignment_for_source(payee_rules, source_name)
        if not assignment:
            continue
        proxy = proxy_by_source.get(source_key)
        manual_alias = manual_aliases.get(source_key)
        expected_name = proxy["proxy_name"] if proxy else (manual_alias["platform_name"] if manual_alias else source_name)
        expected_phone = (proxy.get("proxy_phone") if proxy else (manual_alias.get("phone") if manual_alias else assignment.get("source_phone"))) or ""
        expected.append({
            "payee_id": f"{'proxy' if proxy else 'direct'}:{source_name}:{proxy['proxy_name'] if proxy else source_name}",
            "name": expected_name,
            "phone": expected_phone,
            "type": assignment.get("type") or "玩家",
            "expected_amount": assignment.get("amount"),
            "proxy_for": source_name if proxy else "",
        })
    # 人工保留优先于 JSON：该人员不应再被识别为 EPC 多余人员或被 JSON 覆盖。
    manual_kept = payee_rules.get("_manual_kept") or []
    manual_keys = {_normalize_person_name(item.get("name")) for item in manual_kept}
    expected = [item for item in expected if _normalize_person_name(item.get("name")) not in manual_keys]
    for item in manual_kept:
        name = item.get("name") or ""
        if not _normalize_person_name(name):
            continue
        expected.append({
            "payee_id": f"manual_keep:{name}",
            "name": name,
            "phone": item.get("phone") or "",
            "type": item.get("type") or "玩家",
            "expected_amount": item.get("amount"),
            "proxy_for": "",
            "manual_keep": True,
        })
    return expected


def verify_actual_payees(
    payee_rules: dict,
    actual_rows: list[dict],
    expected_grand_total: float | int | None = None,
) -> list[dict]:
    """新版核验：金额/类型/代收/支付宝字段/页面红色错误全部纳入结果。"""
    expected = _expected_payees_from_rules(payee_rules)
    unused = set(range(len(actual_rows)))
    results: list[dict] = []
    for item in expected:
        candidates = []
        if item["phone"]:
            candidates = [index for index in unused if actual_rows[index].get("phone") == item["phone"]]
        if not candidates and item["name"]:
            candidates = [
                index for index in unused
                if _normalize_person_name(actual_rows[index].get("name")) == _normalize_person_name(item["name"])
            ]
        if not candidates:
            results.append({**item, "actual_amount": None, "result": "missing_in_epc"})
            continue
        index = candidates[0]
        unused.remove(index)
        actual = actual_rows[index]
        expected_amount = _amount_value(item.get("expected_amount"))
        actual_amount = actual.get("actual_amount")
        if actual.get("validation_errors"):
            result = "platform_validation_error"
        elif item.get("phone") and actual.get("phone") != item.get("phone"):
            result = "phone_mismatch"
        elif item.get("type") == "玩家" and actual.get("alipay_login") == "":
            result = "alipay_login_missing"
        elif item.get("type") == "玩家" and actual.get("alipay_name") == "":
            result = "alipay_real_name_missing"
        elif len(candidates) > 1:
            result = "duplicate_in_epc"
        elif item.get("type") and item["type"] not in (actual.get("type") or ""):
            result = "type_mismatch"
        elif expected_amount is None or actual_amount is None or abs(expected_amount - actual_amount) > 0.01:
            result = "amount_mismatch"
        else:
            result = "matched"
        results.append({
            **item,
            "actual_amount": actual_amount,
            "actual_name": actual.get("name", ""),
            "actual_phone": actual.get("phone", ""),
            "actual_type": actual.get("type", ""),
            "actual_alipay_login": actual.get("alipay_login"),
            "actual_alipay_name": actual.get("alipay_name"),
            "platform_errors": actual.get("validation_errors", []),
            "page": actual.get("page"),
            "result": result,
        })
    for index in sorted(unused):
        actual = actual_rows[index]
        results.append({
            "payee_id": "",
            "name": actual.get("name", ""),
            "phone": actual.get("phone", ""),
            "type": actual.get("type", ""),
            "expected_amount": None,
            "actual_amount": actual.get("actual_amount"),
            "actual_alipay_login": actual.get("alipay_login"),
            "actual_alipay_name": actual.get("alipay_name"),
            "platform_errors": actual.get("validation_errors", []),
            "page": actual.get("page"),
            "result": "extra_in_epc",
        })
    if expected_grand_total is not None:
        expected_total = _amount_value(expected_grand_total)
        actual_total = sum(_amount_value(row.get("actual_amount")) or 0 for row in actual_rows)
        if expected_total is not None:
            results.append({
                "payee_id": "__total__",
                "name": "收款人实际合计",
                "phone": "",
                "type": "总金额对账",
                "expected_amount": expected_total,
                "actual_amount": actual_total,
                "page": 0,
                "result": "matched" if abs(expected_total - actual_total) <= 0.01 else "total_amount_mismatch",
            })
    return results


async def run_order_in_page(page: Page, data: dict, runtime: Any) -> str:
    """在批次 Worker 分配的独立 EPC 标签页内执行一张标准单据。"""
    page1_data, expense_note = _page1_payload(data)
    project_id = data.get("project_id") or ""

    duplicate = page.locator('[class*="theoryBox"]')
    if await duplicate.count() > 0:
        duplicate_text = (await duplicate.inner_text())[:300]
        decision = (await _ask_runtime(
            runtime,
            f"项目 {project_id} 可能已有报销单：{duplicate_text}。确认继续提报？[y/n]",
            "WAITING_PRECHECK_APPROVAL",
        )).strip().lower()
        if decision != "y":
            runtime.log("用户取消继续重复项目")
            return "requires_attention"

    await _checkpoint(runtime, "第1页填写前")
    runtime.set_status("RUNNING", "填写第1页报销内容")
    await fill_page1(page, page1_data, expense_note, runtime=runtime)
    await runtime.screenshot(page, "page1")

    decision = (await _ask_runtime(
        runtime,
        "第1页已填写并截图。请检查内容及产品确认截图；输入 y 由脚本进入收款人页，输入 n 停止。",
        "WAITING_PAGE1_APPROVAL",
    )).strip().lower()
    if decision != "y":
        runtime.log("第1页未获继续确认")
        return "requires_attention"

    await _checkpoint(runtime, "第1页确认后")
    advanced = await wait_and_click_next(
        page,
        runtime,
        "第1页",
        'th:has-text("真实姓名")',
    )
    if not advanced:
        return "requires_attention"
    await page.wait_for_timeout(600)

    runtime.set_status("FILLING_PAYEES", "读取并核对 EPC 收款人名单")
    fill_result = await fill_page2_by_rules(page, data.get("payee_rules", {}), runtime=runtime)
    if not fill_result or fill_result.get("stopped"):
        runtime.log("收款人填写被停止")
        return "requires_attention"

    await _checkpoint(runtime, "收款人填写后")
    runtime.set_status("VERIFYING_PAYEES", "回读 EPC 收款人金额")
    auto_repair_attempts = 0
    while True:
        actual_rows = await read_actual_payees(page)
        verification = verify_actual_payees(
            data.get("payee_rules", {}),
            actual_rows,
            expected_grand_total=data.get("grand_total"),
        )
        runtime.save_verification(verification)
        raw_failures = [item for item in verification if item.get("result") != "matched"]
        ignored_issue_keys = set(data.get("payee_rules", {}).get("_ignored_verification_issues") or [])
        ignored_failures = [item for item in raw_failures if _verification_issue_key(item) in ignored_issue_keys]
        failures = [item for item in raw_failures if _verification_issue_key(item) not in ignored_issue_keys]
        if ignored_failures:
            runtime.log("↷ 已忽略本单核验项：" + " | ".join(verification_issue_messages(ignored_failures)))
        total_only_from_manual_keep = (
            bool(data.get("payee_rules", {}).get("_manual_kept"))
            and all(item.get("result") == "total_amount_mismatch" for item in failures)
        )
        if total_only_from_manual_keep and data.get("payee_rules", {}).get("_manual_total_accepted"):
            runtime.log("✓ 已接受人工保留导致的总金额差额，继续完成当前单")
            break
        if not failures:
            break
        auto_corrections = automatic_verification_corrections(failures)
        if auto_corrections and auto_repair_attempts < 2:
            auto_repair_attempts += 1
            runtime.log(
                "↩ JSON 自动回滚：" + " | ".join(
                    f"{item.get('actual_name') or item.get('name')}→{'删除' if item.get('action') == 'delete' else '恢复草稿身份/金额'}"
                    for item in auto_corrections
                )
            )
            auto_errors = await apply_verification_corrections(page, auto_corrections)
            if not auto_errors:
                runtime.log("✓ JSON 自动回滚已写入 EPC，开始重新核验")
                continue
            runtime.log("⚠ JSON 自动回滚未完全成功：" + " | ".join(auto_errors))
        issue_messages = verification_issue_messages(failures)
        runtime.log("⚠ 收款人核验异常：" + " | ".join(issue_messages))
        review_payload = verification_review_payload(failures)
        review_payload["allow_manual_total_acceptance"] = total_only_from_manual_keep
        print("__PAYEE_REVIEW__" + json.dumps(review_payload, ensure_ascii=False))
        issue_summary = "；".join(issue_messages)
        decision = (await _ask_runtime(
            runtime,
            f"收款人核验异常：{issue_summary}。请在前端差异表修改后点击“应用修正并重新核验”；也可在 EPC 当前页手动修正后点击“已修正，重新核验”。",
            "REQUIRES_ATTENTION",
        )).strip().lower()
        if decision.startswith("{"):
            try:
                correction_payload = json.loads(decision)
                corrections = correction_payload.get("verification_corrections") or []
                requested_ignored_keys = {
                    str(key) for key in (correction_payload.get("ignored_issue_keys") or []) if str(key)
                }
                valid_issue_keys = {_verification_issue_key(item) for item in raw_failures}
                new_ignored_keys = requested_ignored_keys & valid_issue_keys
                if new_ignored_keys:
                    rules = data.setdefault("payee_rules", {})
                    existing_ignored = set(rules.get("_ignored_verification_issues") or [])
                    rules["_ignored_verification_issues"] = sorted(existing_ignored | new_ignored_keys)
                    if runtime is not None:
                        runtime.store.update_order(runtime.order_id, draft=runtime.data)
                    ignored_names = [
                        item.get("name") or item.get("actual_name") or "未命名人员"
                        for item in raw_failures if _verification_issue_key(item) in new_ignored_keys
                    ]
                    runtime.log("↷ 已记录忽略核验项：" + "、".join(ignored_names))
                if correction_payload.get("accept_total_mismatch") and total_only_from_manual_keep:
                    data.setdefault("payee_rules", {})["_manual_total_accepted"] = True
                    if runtime is not None:
                        runtime.store.update_order(runtime.order_id, draft=runtime.data)
                    runtime.log("✓ 已确认保留人员优先，总金额差额仅记录不回滚")
                    continue
                apply_errors = await apply_verification_corrections(page, corrections)
                if apply_errors:
                    runtime.log("⚠ 前端修正未完全写入：" + " | ".join(apply_errors))
                    print("__PAYEE_REVIEW__" + json.dumps({
                        "verification_issues": failures,
                        "apply_errors": apply_errors,
                    }, ensure_ascii=False))
                    continue
                runtime.log("✓ 已按前端差异表回写 EPC，开始重新核验")
                continue
            except Exception as exc:
                runtime.log(f"⚠ 无法处理前端修正：{exc}")
                continue
        if decision != "y":
            return "requires_attention"

    await runtime.screenshot(page, "page2_verified")
    decision = (await _ask_runtime(
        runtime,
        "收款人填写及金额核验已通过。请点击“确认核验无误，完成提报”；系统将本单标记为提报完成，但不会进入第3页或点击最终提交。",
        "WAITING_PAGE2_APPROVAL",
    )).strip().lower()
    if decision != "y":
        return "requires_attention"
    runtime.log("✓ 第2页已确认：本单填写完成（未进入第3页，未点击最终提交）")
    return "completed"


async def run(excel_path: str, payee_excel_path: str = None):
    from parser import parse_excel, build_expense_note
    import openpyxl

    print(f"\n{'='*50}")
    print("EPC 报销自动化 Agent")
    print(f"{'='*50}")

    # 解析 Excel
    print(f"\n读取输入文件: {excel_path}")
    data = parse_excel(excel_path)

    ov = data["overview"]
    print(f"  项目: {ov['project_name']} ({ov['project_id']})")
    print(f"  报销明细: {ov['reimbursement_types']}")
    print(f"  收款人: {len(data['payees'])} 人")

    expense_note = build_expense_note(data)
    print(f"\n费用说明（预览）:\n  {expense_note[:100]}...")

    # 生成收款人上传文件
    if payee_excel_path is None:
        payee_excel_path = excel_path.replace(".xlsx", "_收款人上传.xlsx")
        _gen_payee_upload(data["payees"], payee_excel_path)
        print(f"\n已生成收款人上传文件: {payee_excel_path}")

    # 确认开始
    print("\n即将打开浏览器，请确认以上信息正确...")
    cmd = input("输入 'y' 开始自动化，其他键退出: ").strip().lower()
    if cmd != "y":
        print("已退出")
        return

    # 启动浏览器
    async with async_playwright() as pw:
        if CHROME_USER_DATA:
            browser = await pw.chromium.launch_persistent_context(
                user_data_dir=CHROME_USER_DATA,
                channel=BROWSER_CHANNEL or None,
                headless=False,
                slow_mo=SLOW_MO,
                args=["--start-maximized"],
            )
            page = browser.pages[0] if browser.pages else await browser.new_page()
        else:
            browser = await pw.chromium.launch(
                channel=BROWSER_CHANNEL or None,
                headless=False,
                slow_mo=SLOW_MO,
                args=["--start-maximized"],
            )
            context = await browser.new_context(viewport={"width": 1920, "height": 1080})
            page = await context.new_page()

        try:
            # 登录
            await login(page)

            # 导航到报销页面
            print("\n打开报销页面...")
            await page.goto(EPC_URL, wait_until="networkidle")
            await page.wait_for_timeout(2000)

            # 检查重复报销警告
            duplicate_warning = page.locator('[class*="theoryBox"]')
            if await duplicate_warning.count() > 0:
                print("\n⚠ 警告：该项目已存在报销单，请确认是否为拆单！")
                text = await duplicate_warning.text_content()
                print(f"  {text[:200]}")
                cmd = input("  确认继续？[y/n]: ").strip().lower()
                if cmd != "y":
                    print("已取消")
                    return

            # 第1页
            await fill_page1(page, data, expense_note)

            # 截图第1页
            await page.screenshot(path="page1_filled.png", full_page=True)
            print("\n📸 第1页截图: page1_filled.png")
            cmd = input("确认第1页无误后输入 'y' 继续: ").strip().lower()
            if cmd != "y":
                print("已暂停，请手动检查")
                input("按 Enter 退出...")
                return

            # 等 EPC 校验完成后自动点【下一步】（不受固定 30 秒限制）
            await wait_and_click_next(page, None, "第1页", 'th:has-text("真实姓名")')
            await page.wait_for_timeout(2000)
            print("\n→ 进入第2页")

            # 第2页
            await fill_page2(page, payee_excel_path)

            # 截图第2页
            await page.screenshot(path="page2_filled.png", full_page=True)
            print("📸 第2页截图: page2_filled.png")
            cmd = input("确认第2页无误后输入 'y' 继续: ").strip().lower()
            if cmd != "y":
                print("已暂停")
                input("按 Enter 退出...")
                return

            # 等 EPC 校验完成后自动点下一步进入第3页
            await wait_and_click_next(page, None, "收款人页", 'button:has-text("提交")')
            await page.wait_for_timeout(2000)
            print("\n→ 进入第3页")

            # 第3页
            await handle_page3(page, data.get("split_note", {}))

        finally:
            input("\n按 Enter 关闭浏览器...")
            await browser.close()


def _gen_payee_upload(payees: list, out_path: str):
    """生成符合平台批量上传格式的 Excel"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "收款人"

    headers = [
        "玩家姓名", "收款人类型", "联系方式",
        "性别", "国家(地区)", "出生日期",
        "证件类型", "证件号码",
        "支付宝登录号", "支付宝实名姓名", "礼金金额(元)"
    ]
    ws.append(headers)

    for p in payees:
        row = [
            p.get("玩家姓名", ""),
            p.get("收款人类型") or "玩家",
            p.get("联系方式", ""),
            p.get("性别", ""),
            p.get("国家(地区)", ""),
            p.get("出生日期", ""),
            p.get("证件类型", ""),
            str(p.get("证件号码", "")),
            str(p.get("支付宝登录号", "")),
            p.get("支付宝实名姓名", ""),
            p.get("礼金金额(元)", ""),
        ]
        ws.append(row)

    # 设置文本格式防止身份证科学计数法
    from openpyxl.styles import numbers
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"

    wb.save(out_path)


if __name__ == "__main__":
    import sys
    excel = sys.argv[1] if len(sys.argv) > 1 else "报销提报输入模板_v1.xlsx"
    asyncio.run(run(excel))
